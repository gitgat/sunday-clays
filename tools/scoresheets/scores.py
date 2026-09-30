"""Official scores from the cumulative scores workbook (the "ALL SCORE DETAIL" sheet), grouped by Sunday."""

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

SHEET = "ALL SCORE DETAIL"
HEADER = ("Name", "Score Shot", "Event")
HEADER_SEARCH_ROWS = 20
MAX_SCORE = 50


class ScoresError(Exception):
    """The workbook cannot be used as the source of official scores."""


def _label(value: object) -> str:
    return re.sub(r"\s+", "", str(value or "")).lower()


def _score(value: object) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, str) and value.strip().isdigit():
        value = int(value.strip())
    if isinstance(value, int) and 0 <= value <= MAX_SCORE:
        return value
    return None


def _day(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    return value.isoformat() if isinstance(value, date) else None


def load_officials(path: Path) -> dict[str, list[dict[str, Any]]]:
    """{ISO date: rows in workbook order}, each {"name" as written, "hits", "note", "gauge"}.

    A shooter with two rounds on a Sunday has two rows. Rows without a name, a 0-50 whole score or a date are left out.
    """
    try:
        book = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # unreadable, missing or not a workbook: the reader raises many kinds
        raise ScoresError(f"cannot read the scores workbook {path}: {exc}") from exc
    try:
        sheet = next((s for s in book.worksheets if _label(s.title) == _label(SHEET)), None)
        if sheet is None:
            raise ScoresError(f"the scores workbook has no '{SHEET}' sheet")
        rows = sheet.iter_rows(values_only=True)
        columns: dict[str, int] = {}
        for _, row in zip(range(HEADER_SEARCH_ROWS), rows, strict=False):
            found: dict[str, int] = {}
            for index, cell in enumerate(row):
                for wanted in HEADER:
                    if _label(cell) == _label(wanted) and wanted not in found:
                        found[wanted] = index
            if len(found) == len(HEADER):
                columns = found
                break
        if not columns:
            raise ScoresError(f"the '{SHEET}' sheet has no header row with the columns {', '.join(HEADER)}")
        found_rows: dict[str, list[dict[str, Any]]] = {}
        for row in rows:
            cells = {label: row[index] if index < len(row) else None for label, index in columns.items()}
            name = str(cells["Name"] or "").strip()
            hits = _score(cells["Score Shot"])
            iso = _day(cells["Event"])
            if name and hits is not None and iso is not None:
                found_rows.setdefault(iso, []).append({"name": name, "hits": hits, "note": None, "gauge": None})
        return found_rows
    finally:
        book.close()
