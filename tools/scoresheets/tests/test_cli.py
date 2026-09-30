import json
import runpy
import sys
from pathlib import Path
from typing import ClassVar

import pytest
from PIL import Image

import cli
import export
import pipeline
from tests.fixtures import PDF_TEXT
from tests.helpers import make_out
from vlm import VlmUnreachable

REAL_DEFAULT_SCORES = cli.DEFAULT_SCORES


@pytest.fixture(autouse=True)
def no_default_workbook(monkeypatch, tmp_path_factory):
    """The real backend fixture workbook is slow to load and not what these tests are about."""
    monkeypatch.setattr(cli, "DEFAULT_SCORES", tmp_path_factory.mktemp("none") / "absent.xlsx")


def run(capsys, *argv):
    code = cli.main(list(argv))
    captured = capsys.readouterr()
    return code, captured.out, captured.err


@pytest.fixture
def fake_pdf(tmp_path, monkeypatch):
    monkeypatch.setattr(pipeline, "run_pdftotext", lambda pdf: PDF_TEXT)
    monkeypatch.setattr(pipeline, "render_page", lambda pdf, page: Image.new("RGB", (1700, 2200), "white"))
    pdf = tmp_path / "a.pdf"
    pdf.write_bytes(b"a")
    return pdf


class Model:
    """Reads every block as Ann Testerson with a sheet that matches the official 22."""

    def __init__(self, cache_dir):
        pass

    def ask(self, image, prompt, temperature, attempt=0):
        if "Tot column" in prompt:
            return '{"tots": [7, 7, 8], "event_total": 22}'
        return '{"name": "Ann Testerson", "malf": "", "event_total": 22}'

    def close(self):
        pass


def test_scan_then_read_review_export(tmp_path, fake_pdf, monkeypatch, capsys):
    out = tmp_path / "out"
    monkeypatch.setattr(cli, "VlmClient", Model)
    code, stdout, _ = run(capsys, "--out", str(out), "scan", str(fake_pdf))
    assert code == 0 and "7 results, 3 stations (22 targets), 6 blocks" in stdout and "scanned 1 Sundays" in stdout
    code, stdout, _ = run(capsys, "--out", str(out), "read", "--date", "2026-05-31", "--jobs", "2")
    assert code == 0 and "2026-05-31: 1 ok, 5 review, 0 blank" in stdout
    code, stdout, _ = run(capsys, "--out", str(out), "review")
    assert code == 0 and "5 blocks to review" in stdout
    code, stdout, _ = run(capsys, "--out", str(out), "export")
    assert code == 0 and "1 sheets, 1 shooter rows" in stdout
    assert (out / "stations_backfill.xlsx").exists()


def test_read_ignores_other_dates(tmp_path, capsys, monkeypatch):
    out = make_out(tmp_path)
    monkeypatch.setattr(cli, "VlmClient", Model)
    code, stdout, _ = run(capsys, "--out", str(out), "read", "--date", "2020-01-05")
    assert code == 0 and stdout == ""


def test_read_stops_when_the_server_is_unreachable(tmp_path, capsys, monkeypatch):
    out = make_out(tmp_path)

    class Down(Model):
        def ask(self, *args, **kwargs):
            raise VlmUnreachable("cannot reach the model server")

    monkeypatch.setattr(cli, "VlmClient", Down)
    code, _, err = run(capsys, "--out", str(out), "read")
    assert code == 2 and "cannot reach" in err


def test_apply_review_exit_codes(tmp_path, capsys):
    out = make_out(tmp_path)
    good = tmp_path / "good.json"
    good.write_text(json.dumps({"decisions": [{"date": "2026-05-31", "id": "p6-b3", "action": "skip"}]}))
    code, stdout, _ = run(capsys, "--out", str(out), "apply-review", str(good))
    assert code == 0 and "applied 1 decisions" in stdout
    bad = tmp_path / "bad.json"
    bad.write_text(json.dumps({"decisions": [{"date": "2026-05-31", "id": "nope", "action": "skip"}]}))
    code, _, err = run(capsys, "--out", str(out), "apply-review", str(bad))
    assert code == 1 and "no such block" in err


def test_verify_needs_a_workbook_then_reports(tmp_path, capsys, monkeypatch):
    out = make_out(tmp_path)
    code, _, err = run(capsys, "--out", str(out), "verify")
    assert code == 1 and "run export first" in err
    (out / "stations_backfill.xlsx").write_bytes(b"x")
    monkeypatch.setattr(export, "verify_workbook", lambda path, backend: (True, "1 sheets, 2 rows"))
    assert run(capsys, "--out", str(out), "verify")[:2] == (0, "1 sheets, 2 rows\n")
    monkeypatch.setattr(export, "verify_workbook", lambda path, backend: (False, "ERROR x"))
    assert run(capsys, "--out", str(out), "verify")[0] == 1


def test_run_end_to_end_prints_summary(tmp_path, fake_pdf, monkeypatch, capsys):
    out = tmp_path / "out"
    monkeypatch.setattr(cli, "VlmClient", Model)
    code, stdout, _ = run(capsys, "--out", str(out), "run", str(fake_pdf), "--skip-date", "2020-01-05")
    assert code == 0
    assert "Sundays 1: shooters 1 ok, 5 review, 0 blank" in stdout
    assert "exported 1 sheets" in stdout and str(out / "review.html") in stdout


def test_run_summary_lists_sheet_vs_official(tmp_path, capsys):
    out = make_out(tmp_path)
    cli._summary(out, ["2026-05-31"])
    text = capsys.readouterr().out
    assert "shooters 1 ok, 3 review, 1 blank" in text
    assert "sheet_vs_official 2026-05-31 p6-b2: sheet 17 vs official 21" in text
    cli._summary(out, [])
    assert "Sundays 0: shooters 0 ok" in capsys.readouterr().out


def test_run_stops_when_the_model_is_unreachable(tmp_path, fake_pdf, monkeypatch, capsys):
    class Down(Model):
        def ask(self, *args, **kwargs):
            raise VlmUnreachable("down")

    monkeypatch.setattr(cli, "VlmClient", Down)
    code, _, err = run(capsys, "--out", str(tmp_path / "out"), "run", str(fake_pdf))
    assert code == 2 and "down" in err


def test_log_goes_to_stderr(capsys):
    cli._log("hello")
    assert capsys.readouterr().err == "hello\n"


def test_module_entry_point(monkeypatch):
    monkeypatch.setattr(sys, "argv", ["cli.py", "review", "--out"])
    with pytest.raises(SystemExit) as exc:
        runpy.run_path(str(Path(cli.__file__)), run_name="__main__")
    assert exc.value.code == 2


def test_read_skips_a_sunday_whose_tables_are_pictures(tmp_path, capsys, monkeypatch):
    out = make_out(tmp_path)
    sunday = out / "sundays" / "2026-05-31" / "sunday.json"
    sunday.write_text(json.dumps({**json.loads(sunday.read_text()), "layout": []}))
    monkeypatch.setattr(cli, "VlmClient", Model)
    code, stdout, _ = run(capsys, "--out", str(out), "read")
    assert code == 0 and "no typed results or course" in stdout


def test_export_reports_left_out_sundays(tmp_path, capsys):
    out = make_out(tmp_path)
    sunday = out / "sundays" / "2026-05-31" / "sunday.json"
    meta = json.loads(sunday.read_text())
    sunday.write_text(json.dumps({**meta, "layout": [], "course_page": None}))
    code, stdout, _ = run(capsys, "--out", str(out), "export")
    assert code == 0 and "1 Sundays left out" in stdout and "no_course 1" in stdout
    assert "course_unreadable" not in stdout


def test_export_says_which_shooters_station_totals_miss_the_spreadsheet_score(tmp_path, capsys):
    out = make_out(tmp_path)
    readings = out / "sundays" / "2026-05-31" / "readings.json"
    rows = json.loads(readings.read_text())
    rows[1].update(status="accepted", reasons=[])  # Bo: tots add to 17, the spreadsheet says 21
    readings.write_text(json.dumps(rows))
    code, stdout, err = run(capsys, "--out", str(out), "export")
    assert code == 0 and "wrote 1 sheets, 1 shooter rows" in stdout
    assert "left out: 2026-05-31 p6-b2: station totals add to 17, spreadsheet score is 21" in err


def test_export_names_every_reason_a_sunday_was_left_out(tmp_path, capsys):
    out = make_out(tmp_path)
    sunday = out / "sundays" / "2026-05-31" / "sunday.json"
    meta = json.loads(sunday.read_text())
    sunday.write_text(json.dumps({**meta, "layout": [], "course_unreadable": True}))
    other = out / "sundays" / "2026-06-07"
    other.mkdir()
    (other / "sunday.json").write_text(json.dumps({**meta, "date": "2026-06-07", "layout": [], "course_page": None}))
    code, stdout, _ = run(capsys, "--out", str(out), "export")
    assert code == 0 and "2 Sundays left out" in stdout
    assert "course_unreadable 1" in stdout and "no_course 1" in stdout


def test_read_reports_kept_review_decisions(tmp_path, capsys, monkeypatch):
    out = make_out(tmp_path)
    readings = out / "sundays" / "2026-05-31" / "readings.json"
    decided = json.loads(readings.read_text())
    decided[1]["status"] = "accepted"
    readings.write_text(json.dumps(decided))
    sunday = out / "sundays" / "2026-05-31" / "sunday.json"
    meta = json.loads(sunday.read_text())
    sunday.write_text(json.dumps({**meta, "blocks": [r["id"] for r in decided]}))  # every block has an image
    monkeypatch.setattr(cli, "VlmClient", Model)
    code, stdout, _ = run(capsys, "--out", str(out), "read")
    assert code == 0 and "1 review decisions kept" in stdout


# --- follow-ups --------------------------------------------------------------------------------------------------


def workbook_file(tmp_path):
    from datetime import datetime

    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "ALL SCORE DETAIL"
    sheet.append(["Name", "Score Shot", "Event"])
    sheet.append(["Testerson, Ann", 22, datetime(2026, 5, 31)])
    sheet.append(["Fakeman, Bo", 21, datetime(2026, 5, 31)])
    path = tmp_path / "scores.xlsx"
    book.save(path)
    return path


def test_scan_takes_official_scores_from_the_workbook(tmp_path, fake_pdf, capsys):
    out = tmp_path / "out"
    code, stdout, _ = run(capsys, "--out", str(out), "--scores", str(workbook_file(tmp_path)), "scan", str(fake_pdf))
    assert code == 0 and "2 results (from the scores workbook)" in stdout
    meta = json.loads((out / "sundays" / "2026-05-31" / "sunday.json").read_text())
    assert [o["name"] for o in meta["officials"]] == ["Testerson, Ann", "Fakeman, Bo"]


def test_run_uses_the_workbook_too(tmp_path, fake_pdf, monkeypatch, capsys):
    monkeypatch.setattr(cli, "VlmClient", Model)
    out = tmp_path / "out"
    code, stdout, _ = run(capsys, "--out", str(out), "--scores", str(workbook_file(tmp_path)), "run", str(fake_pdf))
    assert code == 0 and "Sundays 1: shooters 1 ok, 5 review" in stdout


def test_a_scores_workbook_that_cannot_be_read_stops_the_command(tmp_path, fake_pdf, capsys):
    code, _, err = run(
        capsys, "--out", str(tmp_path / "out"), "--scores", str(tmp_path / "nope.xlsx"), "scan", str(fake_pdf)
    )
    assert code == 2 and "cannot read" in err
    code, _, err = run(
        capsys, "--out", str(tmp_path / "out"), "--scores", str(tmp_path / "nope.xlsx"), "run", str(fake_pdf)
    )
    assert code == 2 and "cannot read" in err


def test_the_default_scores_workbook_is_used_only_when_it_exists(tmp_path, monkeypatch):
    missing = tmp_path / "none.xlsx"
    monkeypatch.setattr(cli, "DEFAULT_SCORES", missing)
    args = cli.build_parser().parse_args(["scan", "a.pdf"])
    assert cli._officials(args) is None
    monkeypatch.setattr(cli, "DEFAULT_SCORES", workbook_file(tmp_path))
    assert list(cli._officials(cli.build_parser().parse_args(["scan", "a.pdf"]))) == ["2026-05-31"]  # type: ignore[arg-type]


def test_default_scores_path_points_at_the_backend_fixture():
    assert REAL_DEFAULT_SCORES.as_posix().endswith("backend/tests/fixtures/scores_2026-09-27.xlsx")


def test_skipped_default_date_is_reported_by_scan(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(pipeline, "run_pdftotext", lambda pdf: "nothing")
    pdf = tmp_path / "2025-11-16 Sunday Clays Update.pdf"
    pdf.write_bytes(b"a")
    code, stdout, _ = run(capsys, "--out", str(tmp_path / "out"), "scan", str(pdf))
    assert code == 0 and "no scores" in stdout and "scanned 0 Sundays" in stdout


def image_course_out(tmp_path, **extra):
    out = make_out(tmp_path)
    sunday = out / "sundays" / "2026-05-31" / "sunday.json"
    meta = json.loads(sunday.read_text())
    readings = json.loads((out / "sundays" / "2026-05-31" / "readings.json").read_text())
    blocks = [r["id"] for r in readings]  # every block has an image
    sunday.write_text(json.dumps({**meta, "layout": [], "course_page": 2, "blocks": blocks, **extra}))
    Image.new("RGB", (4, 4), "white").save(out / "sundays" / "2026-05-31" / "course.png")
    return out, sunday


class CourseModel(Model):
    def ask(self, image, prompt, temperature, attempt=0):
        if "stations" in prompt:
            return '{"stations": [{"stn": 2, "targets": 7}, {"stn": 4, "targets": 7}, {"stn": 7, "targets": 36}]}'
        return super().ask(image, prompt, temperature, attempt)


def test_read_takes_the_course_from_the_picture_first(tmp_path, capsys, monkeypatch):
    out, sunday = image_course_out(tmp_path)
    monkeypatch.setattr(cli, "VlmClient", CourseModel)
    code, stdout, _ = run(capsys, "--out", str(out), "read")
    assert code == 0 and "course read from the picture" in stdout
    assert json.loads(sunday.read_text())["layout"] == [[2, 7], [4, 7], [7, 36]]
    assert "2026-05-31: " in stdout and "review" in stdout


def test_read_skips_a_sunday_whose_course_picture_cannot_be_read(tmp_path, capsys, monkeypatch):
    out, sunday = image_course_out(tmp_path)

    class Bad(Model):
        def ask(self, *args, **kwargs):
            return "garbage"

    monkeypatch.setattr(cli, "VlmClient", Bad)
    code, stdout, _ = run(capsys, "--out", str(out), "read")
    assert code == 0 and "course_unreadable" in stdout
    assert json.loads(sunday.read_text())["course_unreadable"] is True
    assert "not read" in stdout


def test_read_stops_when_the_model_is_unreachable_during_the_course_read(tmp_path, capsys, monkeypatch):
    out, _ = image_course_out(tmp_path)

    class Down(Model):
        def ask(self, *args, **kwargs):
            raise VlmUnreachable("down")

    monkeypatch.setattr(cli, "VlmClient", Down)
    code, _, err = run(capsys, "--out", str(out), "read")
    assert code == 2 and "down" in err


def test_a_sunday_with_a_lettered_station_is_scanned_read_and_exported(tmp_path, monkeypatch, capsys):
    typed = PDF_TEXT.replace(
        "                        7    8     1   SINGLE", "                        7A   8     1   SINGLE"
    )
    assert typed != PDF_TEXT
    monkeypatch.setattr(pipeline, "run_pdftotext", lambda pdf: typed)
    monkeypatch.setattr(pipeline, "render_page", lambda pdf, page: Image.new("RGB", (1700, 2200), "white"))
    pdf = tmp_path / "a.pdf"
    pdf.write_bytes(b"a")
    out = tmp_path / "out"
    monkeypatch.setattr(cli, "VlmClient", Model)
    code, stdout, _ = run(capsys, "--out", str(out), "scan", str(pdf))
    assert code == 0 and "3 stations (22 targets)" in stdout
    code, stdout, _ = run(capsys, "--out", str(out), "read")
    assert code == 0 and "1 ok, 5 review" in stdout and "not read" not in stdout
    code, stdout, _ = run(capsys, "--out", str(out), "export")
    assert code == 0 and "1 sheets, 1 shooter rows" in stdout and "left out" not in stdout
    from openpyxl import load_workbook

    sheet = load_workbook(out / "stations_backfill.xlsx")["5 31 26"]
    assert [c.value for c in sheet[4]][:4] == ["STATION #", 2, 4, "7A"]


def test_problems_are_printed_to_stderr(capsys):
    cli._print_problems(["2026-05-31 p6-b2: station totals add to 17, spreadsheet score is 21"])
    assert capsys.readouterr().err == "left out: 2026-05-31 p6-b2: station totals add to 17, spreadsheet score is 21\n"


def test_read_reports_stale_decisions_dropped(tmp_path, capsys, monkeypatch):
    out = make_out(tmp_path)
    sunday = out / "sundays" / "2026-05-31" / "sunday.json"
    readings = out / "sundays" / "2026-05-31" / "readings.json"
    decided = json.loads(readings.read_text())
    decided[1].update(status="accepted", pdf_sha256="old")
    readings.write_text(json.dumps(decided))
    sunday.write_text(
        json.dumps({**json.loads(sunday.read_text()), "sha256": "new", "blocks": [r["id"] for r in decided]})
    )
    monkeypatch.setattr(cli, "VlmClient", Model)
    code, stdout, _ = run(capsys, "--out", str(out), "read")
    assert code == 0 and "1 stale review decisions dropped (the PDF changed)" in stdout


# --- --engine claude -------------------------------------------------------------------------------------------


class FakeClaude:
    """A ClaudeClient stand-in: every page is Ann Testerson at slot 1 with a sheet matching the official 22."""

    instances: ClassVar[list] = []

    def __init__(self, cache_dir, *, model):
        self.cache_dir, self.model = cache_dir, model
        self.calls, self.cached, self.cost, self.closed = 2, 1, 0.0521, False
        FakeClaude.instances.append(self)

    def ask(self, image, prompt, temperature, attempt=0):
        block = {
            "slot": 1,
            "name": "Ann Testerson",
            "stations": [{"stn": 2, "tot": 7}, {"stn": 4, "tot": 7}],
            "event_total": 22,
        }
        block["stations"].append({"stn": 7, "tot": 8})
        return json.dumps({"blocks": [block]})

    def close(self):
        self.closed = True


@pytest.fixture
def claude(monkeypatch):
    FakeClaude.instances = []
    monkeypatch.setattr(cli.claude_engine, "ClaudeClient", FakeClaude)
    return FakeClaude


def with_pages(out, iso="2026-05-31", page="p6"):
    pages = out / "sundays" / iso / "pages"
    pages.mkdir(exist_ok=True)
    Image.new("RGB", (4, 4), "white").save(pages / f"{page}.png")


def test_engine_claude_reads_pages_and_reports_the_cost(tmp_path, capsys, claude):
    out = make_out(tmp_path)
    with_pages(out)
    code, stdout, _ = run(capsys, "--out", str(out), "--engine", "claude", "--claude-model", "opus", "read")
    assert code == 0 and "2026-05-31: 1 ok, 0 review, 4 blank" in stdout
    assert "claude: 2 calls, 1 cached, cost $0.0521" in stdout
    (client,) = claude.instances
    assert client.model == "opus" and client.cache_dir == out / "claude-cache" and client.closed


def test_engine_claude_uses_a_gemma_cache_only_when_there_is_one(tmp_path, capsys, claude, monkeypatch):
    seen = []

    class Recording(Model):
        def __init__(self, cache_dir, *, offline=False):
            seen.append((cache_dir.name, offline))
            self.closed = False

        def close(self):
            seen.append("closed")

    monkeypatch.setattr(cli, "VlmClient", Recording)
    out = make_out(tmp_path)
    with_pages(out)
    run(capsys, "--out", str(out), "--engine", "claude", "read")
    assert seen == []  # no out/cache: no second opinion, no Gemma client
    (out / "cache").mkdir()
    run(capsys, "--out", str(out), "--engine", "claude", "read")
    assert seen == [("cache", True), "closed"]


def test_jobs_default_to_four_for_claude_and_one_for_gemma_and_take_either_position(
    tmp_path, capsys, claude, monkeypatch
):
    jobs = []
    real = pipeline.read_sunday

    def spy(root, client, *, jobs=1, **kwargs):
        seen_jobs.append(jobs)
        return real(root, client, jobs=jobs, **kwargs)

    seen_jobs = jobs
    monkeypatch.setattr(pipeline, "read_sunday", spy)
    out = make_out(tmp_path)
    with_pages(out)
    run(capsys, "--out", str(out), "--engine", "claude", "read")
    run(capsys, "--out", str(out), "--engine", "claude", "--jobs", "2", "read")
    run(capsys, "--out", str(out), "--engine", "claude", "read", "--jobs", "3")
    for suffix in ("", "-tot"):  # the Gemma engine reads block images, which make_out leaves out for a blank block
        Image.new("RGB", (4, 4), "white").save(out / "sundays" / "2026-05-31" / "blocks" / f"p6-b5{suffix}.png")
    monkeypatch.setattr(cli, "VlmClient", Model)
    run(capsys, "--out", str(out), "read")
    assert jobs == [4, 2, 3, 1]


def test_engine_claude_stops_when_pages_are_missing_or_claude_is_unavailable(tmp_path, capsys, claude, monkeypatch):
    out = make_out(tmp_path)
    code, _, err = run(capsys, "--out", str(out), "--engine", "claude", "read")
    assert code == 2 and "run scan again" in err
    with_pages(out)

    def down(self, *args, **kwargs):
        raise cli.claude_engine.ClaudeUnavailable("cannot run claude")

    monkeypatch.setattr(FakeClaude, "ask", down)
    code, _, err = run(capsys, "--out", str(out), "--engine", "claude", "read")
    assert code == 2 and "cannot run claude" in err


def test_engine_claude_reads_the_course_picture_with_claude(tmp_path, capsys, claude):
    out, sunday = image_course_out(tmp_path)
    with_pages(out)

    def ask(self, image, prompt, temperature, attempt=0):
        if "Tgts column" in prompt:
            return '{"stations": [{"stn": 2, "targets": 7}, {"stn": 4, "targets": 7}, {"stn": 7, "targets": 36}]}'
        return '{"blocks": []}'

    claude.ask = ask
    code, stdout, _ = run(capsys, "--out", str(out), "--engine", "claude", "read")
    assert code == 0 and "course read from the picture" in stdout
    assert json.loads(sunday.read_text())["course_source"] == "image"


def test_engine_defaults_to_gemma():
    args = cli.build_parser().parse_args(["read"])
    assert (args.engine, args.claude_model, args.jobs) == ("gemma", "opus", None)
