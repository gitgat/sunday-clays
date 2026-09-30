import csv
import json
import shutil
import subprocess
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

import export
from tests.helpers import make_out, record


def exported_out(tmp_path):
    readings = [
        record(1, "Ann Testerson", "Testerson, Ann", 22, [7, 7, 8], "ok"),
        record(2, "Bo Fakeman", "Fakeman, Bo", 21, [7, 7, 7], "accepted"),
        record(3, "Cy Mocklee", "Mockley, Cy", 20, [7, None], "review", ["tot_missing"]),
        record(4, "Someone", "Mockley, Cy", 20, [1, 1, 1], "skipped"),
    ]
    out = make_out(tmp_path, readings=readings)
    return out


def test_sheet_name_is_month_day_two_digit_year():
    assert export.sheet_name("2026-05-31") == "5 31 26"
    assert export.sheet_name("2025-10-12") == "10 12 25"


def test_workbook_layout_matches_the_station_parser_format(tmp_path):
    out = exported_out(tmp_path)
    assert export.export_all(out) == {
        "sheets": 1,
        "rows": 2,
        "skipped": 0,
        "mismatched": 0,
        "course_unreadable": 0,
        "no_course": 0,
    }
    sheet = load_workbook(out / "stations_backfill.xlsx")["5 31 26"]
    rows = [[c.value for c in row] for row in sheet.iter_rows()]
    assert rows[0][0] == "Event Date"
    assert rows[1][0].date() == date(2026, 5, 31)
    assert rows[3][:4] == ["STATION #", 2, 4, 7]
    assert rows[4][:4] == ["TARGET COUNT", 7, 7, 8]
    assert rows[8][:2] == ["Name", "Station Hits"]
    assert rows[9][:4] == ["Testerson, Ann", 7, 7, 8]
    assert rows[10][:4] == ["Fakeman, Bo", 7, 7, 7]
    assert len(rows) == 11  # review and skipped shooters are left out


def test_report_and_presentations(tmp_path):
    out = exported_out(tmp_path)
    export.export_all(out)
    rows = list(csv.reader((out / "report.csv").read_text().splitlines()))
    assert rows[0] == [
        "date",
        "block",
        "name_read",
        "matched_name",
        "official",
        "sheet_sum",
        "status",
        "reasons",
        "source",
    ]
    assert rows[2] == [
        "2026-05-31",
        "p6-b2",
        "Bo Fakeman",
        "Fakeman, Bo",
        "21",
        "21",
        "accepted",
        "",
        "gemma",
    ]
    assert len(rows) == 5
    sidecar = json.loads((out / "presentations.json").read_text())
    assert sidecar["2026-05-31"]["layout"] == [[2, 7], [4, 7], [7, 8]]
    assert sidecar["2026-05-31"]["presentations"]["2"][0]["presentation"] == "SINGLE"


def test_no_workbook_when_nothing_is_exportable(tmp_path):
    out = make_out(tmp_path, readings=[record(1, "X", None, None, [1], "review", ["no_official_match"])])
    assert export.export_all(out)["sheets"] == 0
    assert not (out / "stations_backfill.xlsx").exists()
    rows = list(csv.reader((out / "report.csv").read_text().splitlines()))
    assert rows[1][3:5] == ["", ""]


def test_verify_workbook_reports_summary_and_errors(tmp_path, monkeypatch):
    seen = []

    def fake_run(cmd, **kwargs):
        seen.append(cmd)
        return subprocess.CompletedProcess(cmd, 0, stdout="noise\n" + json.dumps(RESULT[0]) + "\n", stderr="")

    RESULT = [{"sheets": 2, "rows": 9, "findings": {"info": 1}, "errors": []}]
    monkeypatch.setattr(export.subprocess, "run", fake_run)
    ok, text = export.verify_workbook(tmp_path / "w.xlsx", tmp_path / "backend")
    assert ok and text == "2 sheets, 9 rows, findings {'info': 1}"
    assert seen[0][:4] == ["uv", "run", "--project", str(tmp_path / "backend")]
    RESULT[0]["errors"] = ["hits_out_of_range: bad"]
    ok, text = export.verify_workbook(tmp_path / "w.xlsx", tmp_path / "backend")
    assert not ok and "ERROR hits_out_of_range: bad" in text


def test_verify_workbook_parser_crash(tmp_path, monkeypatch):
    monkeypatch.setattr(export.subprocess, "run", lambda cmd, **kw: subprocess.CompletedProcess(cmd, 1, "", "boom\n"))
    assert export.verify_workbook(tmp_path / "w.xlsx", tmp_path) == (False, "boom")
    monkeypatch.setattr(export.subprocess, "run", lambda cmd, **kw: subprocess.CompletedProcess(cmd, 1, "", ""))
    assert export.verify_workbook(tmp_path / "w.xlsx", tmp_path) == (False, "the backend parser failed")


BACKEND = Path(__file__).resolve().parents[3] / "backend"


@pytest.mark.skipif(
    not BACKEND.is_dir() or shutil.which("uv") is None,
    reason="needs backend/ and uv to run the real station parser",
)
def test_real_backend_parser_accepts_the_workbook(tmp_path, monkeypatch):
    monkeypatch.undo()  # the network guard is irrelevant to a subprocess; keep env untouched
    out = exported_out(tmp_path)
    export.export_all(out)
    ok, text = export.verify_workbook(out / "stations_backfill.xlsx", BACKEND)
    assert ok, text
    assert text.startswith("1 sheets, 2 rows")


def test_a_sunday_with_a_lettered_station_is_exported_with_its_label(tmp_path):
    out = make_out(tmp_path, readings=[record(1, "Ann Testerson", "Testerson, Ann", 22, [7, 7, 8], "ok")])
    sunday = out / "sundays" / "2026-05-31" / "sunday.json"
    meta = json.loads(sunday.read_text())
    meta["layout"] = [[4, 7], [7, 7], ["7A", 8]]
    sunday.write_text(json.dumps(meta))
    counts = export.export_all(out)
    assert (counts["sheets"], counts["rows"], counts["skipped"]) == (1, 1, 0)
    sheet = load_workbook(out / "stations_backfill.xlsx")["5 31 26"]
    rows = [[c.value for c in row] for row in sheet.iter_rows()]
    assert rows[3][:4] == ["STATION #", 4, 7, "7A"]
    assert rows[9][:4] == ["Testerson, Ann", 7, 7, 8]


def test_a_shooter_whose_station_totals_miss_the_spreadsheet_score_is_never_exported(tmp_path):
    readings = [
        record(1, "Ann Testerson", "Testerson, Ann", 22, [7, 7, 8], "ok"),
        record(2, "Bo Fakeman", "Fakeman, Bo", 21, [7, 7, 3], "accepted"),  # an older decision that adds to 17
        record(3, "Cy Mockley", "Mockley, Cy", 20, [7, 7, 6], "review", ["sheet_vs_official"]),
    ]
    out = make_out(tmp_path, readings=readings)
    problems: list[str] = []
    counts = export.export_all(out, problems)
    assert (counts["rows"], counts["mismatched"]) == (1, 1)
    assert problems == ["2026-05-31 p6-b2: station totals add to 17, spreadsheet score is 21"]
    rows = list(csv.reader((out / "report.csv").read_text().splitlines()))
    assert [(row[1], row[6]) for row in rows[1:4]] == [
        ("p6-b1", "ok"),
        ("p6-b2", "sheet_vs_official"),
        ("p6-b3", "review"),
    ]
    assert rows[2][7] == "station totals add to 17, spreadsheet score is 21"
    sheet = load_workbook(out / "stations_backfill.xlsx")["5 31 26"]
    assert [row[0].value for row in sheet.iter_rows(min_row=10)] == ["Testerson, Ann"]
    export.export_all(out)  # the problem list is optional


def test_sum_problem_is_empty_only_when_the_totals_equal_the_spreadsheet_score():
    assert export.sum_problem(22, 22) == ""
    assert export.sum_problem(21, 22) == "station totals add to 21, spreadsheet score is 22"
    assert export.sum_problem(21, "22 or 23") == "station totals add to 21, spreadsheet score is 22 or 23"
    assert export.sum_problem(21, None) == "station totals add to 21, spreadsheet score is missing"


def test_a_stale_workbook_is_removed_when_nothing_is_exportable(tmp_path):
    out = exported_out(tmp_path)
    export.export_all(out)
    assert (out / "stations_backfill.xlsx").exists()
    readings_file = out / "sundays" / "2026-05-31" / "readings.json"
    readings_file.write_text(json.dumps([record(1, "Cy", "Mockley, Cy", 20, [7], "review", ["tot_missing"])]))
    assert export.export_all(out)["sheets"] == 0
    assert not (out / "stations_backfill.xlsx").exists()


def test_presentations_cover_sundays_that_were_scanned_but_not_read(tmp_path):
    out = make_out(tmp_path)
    unread = out / "sundays" / "2026-06-07"
    unread.mkdir(parents=True)
    meta = {"date": "2026-06-07", "layout": [[1, 5]], "presentations": {"1": []}, "officials": [], "blocks": []}
    (unread / "sunday.json").write_text(json.dumps(meta))
    export.export_all(out)
    sidecar = json.loads((out / "presentations.json").read_text())
    assert sorted(sidecar) == ["2026-05-31", "2026-06-07"]
    assert sidecar["2026-06-07"]["layout"] == [[1, 5]]


def test_report_lists_sundays_left_out_with_the_reason(tmp_path):
    out = exported_out(tmp_path)
    sunday = out / "sundays" / "2026-05-31" / "sunday.json"
    meta = json.loads(sunday.read_text())
    sunday.write_text(json.dumps({**meta, "layout": [], "course_page": None}))
    other = out / "sundays" / "2026-06-07"
    other.mkdir()
    (other / "sunday.json").write_text(
        json.dumps({**meta, "date": "2026-06-07", "layout": [], "course_unreadable": True})
    )
    counts = export.export_all(out)
    assert counts == {
        "sheets": 0,
        "rows": 0,
        "skipped": 2,
        "mismatched": 0,
        "course_unreadable": 1,
        "no_course": 1,
    }
    rows = list(csv.reader((out / "report.csv").read_text().splitlines()))
    left_out = [row for row in rows if row[6] == "sunday_skipped"]
    assert [(row[0], row[1], row[7]) for row in left_out] == [
        ("2026-05-31", "", "no_course"),
        ("2026-06-07", "", "course_unreadable"),
    ]


def test_a_sunday_with_no_course_at_all_is_reported_not_dropped_silently(tmp_path):
    out = make_out(tmp_path, readings=[record(1, "Ann Testerson", "Testerson, Ann", 22, [7, 7, 8], "ok")])
    sunday = out / "sundays" / "2026-05-31" / "sunday.json"
    sunday.write_text(json.dumps({**json.loads(sunday.read_text()), "layout": [], "course_page": None}))
    counts = export.export_all(out)
    assert counts["no_course"] == 1 and counts["skipped"] == 1 and counts["sheets"] == 0
    rows = list(csv.reader((out / "report.csv").read_text().splitlines()))
    assert [(row[0], row[6], row[7]) for row in rows if row[6] == "sunday_skipped"] == [
        ("2026-05-31", "sunday_skipped", "no_course")
    ]


def test_readings_from_a_different_pdf_than_the_scan_are_left_out_and_reported_as_stale(tmp_path):
    out = exported_out(tmp_path)
    root = out / "sundays" / "2026-05-31"
    meta = json.loads((root / "sunday.json").read_text())
    (root / "sunday.json").write_text(json.dumps({**meta, "sha256": "new"}))
    readings = json.loads((root / "readings.json").read_text())
    readings[0]["pdf_sha256"] = "old"  # the ok row was read from the previous PDF
    readings[1]["pdf_sha256"] = "new"  # the accepted row matches the scan
    (root / "readings.json").write_text(json.dumps(readings))
    assert export.export_all(out)["rows"] == 1
    rows = list(csv.reader((out / "report.csv").read_text().splitlines()))
    assert [row[6] for row in rows[1:5]] == ["stale", "accepted", "review", "skipped"]
    sheet = load_workbook(out / "stations_backfill.xlsx")["5 31 26"]
    assert [row[0].value for row in sheet.iter_rows(min_row=10)] == ["Fakeman, Bo"]


def test_report_shows_which_engine_gave_each_reading(tmp_path):
    readings = [
        {**record(1, "Ann Testerson", "Testerson, Ann", 22, [7, 7, 8], "ok"), "source": "claude"},
        {**record(2, "Bo Fakeman", "Fakeman, Bo", 21, [7, 7, 7], "ok"), "source": "gemma_second_opinion"},
        record(3, "Cy Mockley", "Mockley, Cy", 20, [7, 7, 6], "ok"),
    ]
    out = make_out(tmp_path, readings=readings)
    export.export_all(out)
    rows = list(csv.reader((out / "report.csv").read_text().splitlines()))
    assert [row[-1] for row in rows[1:4]] == ["claude", "gemma_second_opinion", "gemma"]
