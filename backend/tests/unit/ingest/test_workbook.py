import io
import warnings
import zipfile
from collections.abc import Callable
from datetime import date, datetime, time

import pytest
from openpyxl import Workbook

from sunday_clays.ingest.types import FileKind, ParseError
from sunday_clays.ingest.workbook import (
    MAX_COMPRESSION_RATIO,
    MAX_MEMBER_BYTES,
    MAX_MEMBERS,
    MAX_TOTAL_BYTES,
    LoadedWorkbook,
    _archive_too_large,
    coerce_date,
    coerce_int,
    detect_kind,
    find_sheet,
    is_blank,
    load_workbook_bytes,
    name_text,
    normalize_label,
)

from .builders import loaded_workbook

UNREADABLE = "This file could not be read as an Excel workbook"
TOO_LARGE = "File is too large to process"
MIB = 1024 * 1024
CFB_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")  # OLE2 compound file header


def _zip(members: list[tuple[str, bytes, int]]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, payload, compression in members:
            archive.writestr(name, payload, compress_type=compression)
    return buffer.getvalue()


def _sparse(last_row: int) -> bytes:
    workbook = Workbook()
    workbook.worksheets[0].cell(row=last_row, column=1, value="stray")
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_load_returns_cached_values_and_formula_text() -> None:
    lw = loaded_workbook({"S": [["Doe, Jane", "=1+1"]]})

    assert lw.values["S"]["A1"].value == "Doe, Jane"
    assert lw.values["S"]["B1"].value is None
    assert lw.formulas["S"]["B1"].value == "=1+1"


def test_load_ignores_data_validation_warning(stations_file_bytes: bytes) -> None:
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        load_workbook_bytes(stations_file_bytes)

    assert [str(warning.message) for warning in caught] == []


@pytest.mark.parametrize(
    "make_upload",
    [
        pytest.param(lambda _scores: b"", id="empty"),
        pytest.param(
            lambda _scores: b"Name,Score Shot,Event\nHadley, Ike,34,2026-09-13\n", id="csv-text"
        ),
        pytest.param(lambda scores: scores[:4096], id="truncated-workbook"),
        pytest.param(
            lambda _scores: _zip([("notes.txt", b"not a workbook", zipfile.ZIP_DEFLATED)]),
            id="zip-without-workbook-parts",
        ),
        # Password-protected .xlsx (and legacy .xls) files are OLE2 compound files.
        pytest.param(lambda _scores: CFB_SIGNATURE + bytes(504), id="encrypted-ole2"),
    ],
)
def test_unusable_file_raises_parse_error(
    make_upload: Callable[[bytes], bytes], scores_file_bytes: bytes
) -> None:
    with pytest.raises(ParseError) as excinfo:
        load_workbook_bytes(make_upload(scores_file_bytes))

    # The message is fixed text; the original error is only chained for server logs.
    assert str(excinfo.value) == UNREADABLE
    assert excinfo.value.__cause__ is not None


def test_zip_bomb_member_rejected() -> None:
    bomb = _zip([("xl/worksheets/sheet1.xml", bytes(2 * MIB), zipfile.ZIP_DEFLATED)])

    with pytest.raises(ParseError, match=f"^{TOO_LARGE}$"):
        load_workbook_bytes(bomb)


def test_archive_limits_bound_the_double_load() -> None:
    # The two openpyxl loads peak at ~29x the uncompressed size, so 10 MiB in
    # total keeps one upload under ~300 MB of memory.
    assert (MAX_MEMBERS, MAX_MEMBER_BYTES, MAX_TOTAL_BYTES, MAX_COMPRESSION_RATIO) == (
        500,
        8 * MIB,
        10 * MIB,
        100,
    )


@pytest.mark.parametrize(
    ("members", "message"),
    [
        ([(f"m{i}", b"", zipfile.ZIP_STORED) for i in range(MAX_MEMBERS)], UNREADABLE),
        ([(f"m{i}", b"", zipfile.ZIP_STORED) for i in range(MAX_MEMBERS + 1)], TOO_LARGE),
        ([("big", bytes(MAX_MEMBER_BYTES), zipfile.ZIP_STORED)], UNREADABLE),
        ([("big", bytes(MAX_MEMBER_BYTES + 1), zipfile.ZIP_STORED)], TOO_LARGE),
        (
            [
                ("a", bytes(MAX_MEMBER_BYTES), zipfile.ZIP_STORED),
                ("b", bytes(MAX_TOTAL_BYTES - MAX_MEMBER_BYTES), zipfile.ZIP_STORED),
            ],
            UNREADABLE,
        ),
        (
            [
                ("a", bytes(MAX_MEMBER_BYTES), zipfile.ZIP_STORED),
                ("b", bytes(MAX_TOTAL_BYTES - MAX_MEMBER_BYTES + 1), zipfile.ZIP_STORED),
            ],
            TOO_LARGE,
        ),
    ],
    ids=[
        "max-members",
        "max+1-members",
        "max-member-bytes",
        "max+1-member-bytes",
        "max-total-bytes",
        "max+1-total-bytes",
    ],
)
def test_archive_limits_are_inclusive(members: list[tuple[str, bytes, int]], message: str) -> None:
    # At a limit the archive passes the size check and then fails as a
    # non-workbook; one byte or member over, it is rejected as too large.
    with pytest.raises(ParseError) as excinfo:
        load_workbook_bytes(_zip(members))

    assert str(excinfo.value) == message


def _member(file_size: int, compress_size: int) -> zipfile.ZipInfo:
    info = zipfile.ZipInfo("xl/worksheets/sheet1.xml")
    info.file_size = file_size
    info.compress_size = compress_size
    return info


@pytest.mark.parametrize(
    ("file_size", "compress_size", "too_large"),
    [
        (MAX_COMPRESSION_RATIO * 1000, 1000, False),
        (MAX_COMPRESSION_RATIO * 1000 + 1, 1000, True),
        (MAX_COMPRESSION_RATIO, 0, False),
        (MAX_COMPRESSION_RATIO + 1, 0, True),
        (0, 0, False),
    ],
    ids=[
        "ratio-at-max",
        "ratio-just-over-max",
        "zero-compressed-at-max",
        "zero-compressed-over-max",
        "empty-member",
    ],
)
def test_compression_ratio_limit_is_inclusive(
    file_size: int, compress_size: int, too_large: bool
) -> None:
    # A compress_size of 0 counts as 1 byte, so there is never a division by zero.
    assert _archive_too_large([_member(file_size, compress_size)]) is too_large


def test_sheet_at_the_row_limit_loads() -> None:
    lw = load_workbook_bytes(_sparse(100_000))

    assert lw.values.worksheets[0].max_row == 100_000


def test_sheet_over_the_row_limit_is_rejected() -> None:
    with pytest.raises(ParseError, match=f"^{TOO_LARGE}$"):
        load_workbook_bytes(_sparse(100_001))


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("  STATION   #\n", "station #"),
        ("Name", "name"),
        (4, None),
        (None, None),
    ],
)
def test_normalize_label(value: object, expected: str | None) -> None:
    assert normalize_label(value) == expected


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (datetime.combine(date(2026, 9, 13), time(10, 30)), date(2026, 9, 13)),
        (date(2026, 9, 13), date(2026, 9, 13)),
        (46278, date(2026, 9, 13)),
        (46278.75, date(2026, 9, 13)),
        (61, date(1900, 3, 1)),
        (2_958_465, date(9999, 12, 31)),
        ("9/13/2026", date(2026, 9, 13)),
        (" 2026-09-13 ", date(2026, 9, 13)),
        ("2026-09-13 00:00:00", date(2026, 9, 13)),
        ("2026-09-13T10:30", date(2026, 9, 13)),
        ("2026-09-13T10:30:00Z", date(2026, 9, 13)),
        ("2026-09-13T10:30:00+00:00", date(2026, 9, 13)),
        ("2026-09-13 10:30:00.123456", date(2026, 9, 13)),
        ("2026-09-13T10:30:00.5-07:00", date(2026, 9, 13)),
        ("2026-09-13 whatever", None),
        ("9 13 26", date(2026, 9, 13)),
        ("9 6 26", date(2026, 9, 6)),
        ("2/30/2026", None),
        ("9 13 2026", None),
        ("Sept 13", None),
        ("", None),
        (60, None),
        (2_958_466, None),
        (float("nan"), None),
        (True, None),
        (time(10, 0), None),
        (None, None),
    ],
)
def test_coerce_date(value: object, expected: date | None) -> None:
    assert coerce_date(value) == expected


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (40, 40),
        (0, 0),
        (40.0, 40),
        (" 40 ", 40),
        ("40.0", 40),
        (40.5, None),
        ("forty", None),
        ("", None),
        (float("nan"), None),
        (True, None),
        (False, None),
        (datetime.combine(date(2026, 9, 13), time()), None),
        (None, None),
    ],
)
def test_coerce_int(value: object, expected: int | None) -> None:
    assert coerce_int(value) == expected


@pytest.mark.parametrize(
    ("value", "expected"),
    [(None, True), ("", True), ("  \xa0", True), (0, False), (" x ", False)],
)
def test_is_blank(value: object, expected: bool) -> None:
    assert is_blank(value) is expected


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("Kowalczyk, Barrett\xa0", "Kowalczyk, Barrett\xa0"),
        (123, "123"),
        ("Amos**", "Amos**"),
        ("**", None),
        ("   ", None),
        (None, None),
    ],
)
def test_name_text_keeps_raw_text_or_rejects_keyless_names(
    value: object, expected: str | None
) -> None:
    assert name_text(value) == expected


def test_find_sheet_ignores_case_and_spacing() -> None:
    lw = loaded_workbook({"Attendance  history": [["Date"]]})

    found = find_sheet(lw.values, "Attendance History")

    assert found is not None
    assert found.title == "Attendance  history"
    assert find_sheet(lw.values, "Score Frequency") is None


def test_detect_kind_on_the_real_workbooks(
    scores_lw: LoadedWorkbook, stations_lw: LoadedWorkbook
) -> None:
    assert detect_kind(scores_lw) is FileKind.SCORES
    assert detect_kind(stations_lw) is FileKind.STATIONS


def test_detect_kind_matches_labels_ignoring_case_and_spacing() -> None:
    scores = loaded_workbook({" all  score detail": [["Name"]]})
    stations = loaded_workbook({"9 13 26": [["  event DATE "]]})

    assert detect_kind(scores) is FileKind.SCORES
    assert detect_kind(stations) is FileKind.STATIONS


def test_scores_sheet_wins_when_both_kinds_are_present() -> None:
    both = loaded_workbook({"9 13 26": [["Event Date"]], "ALL SCORE DETAIL": [["Name"]]})

    assert detect_kind(both) is FileKind.SCORES


def test_unrelated_workbook_raises_parse_error() -> None:
    roster = loaded_workbook({"Name List (2)": [["Name", "Class"], ["Hadley, Ike"]]})

    with pytest.raises(ParseError) as excinfo:
        detect_kind(roster)

    assert str(excinfo.value) == (
        "This doesn't look like a Sunday Clays scores or station workbook"
    )
