import hashlib
import io
import json
import warnings
from collections.abc import Mapping
from datetime import date
from pathlib import Path

import openpyxl
import pytest

from sunday_clays.ingest.names import name_key
from sunday_clays.ingest.stations import parse_stations
from sunday_clays.ingest.types import (
    STATION_SHEET_EXCLUSION_CODES,
    Severity,
    StationHitsRow,
    StationLayoutEntry,
    StationsParse,
)
from sunday_clays.ingest.workbook import LoadedWorkbook, load_workbook_bytes

from .builders import SheetRows, loaded_workbook
from .builders_stations import (
    DEVLIN,
    FIRST_ENTRY_ROW,
    HADLEY,
    NEXT_SUNDAY,
    SUNDAY,
    station_sheet,
)

GOLDEN = Path(__file__).resolve().parents[2] / "golden"
FIXTURE_LAYOUT = tuple(
    StationLayoutEntry(station_no, target)
    for station_no, target in zip(range(4, 11), (7, 7, 7, 7, 7, 7, 8), strict=True)
)
HADLEY_HITS = (("4", 3), ("5", 4), ("6", 6), ("7", 5), ("8", 7), ("9", 4), ("10", 7))


@pytest.fixture(scope="module")
def fixture_parse(stations_lw: LoadedWorkbook) -> StationsParse:
    return parse_stations(stations_lw)


def _parse(sheets: Mapping[str, SheetRows]) -> StationsParse:
    return parse_stations(loaded_workbook(sheets))


def _codes(parse: StationsParse) -> list[tuple[str, Severity, str | None, int | None]]:
    return [(f.code, f.severity, f.sheet, f.row) for f in parse.findings]


def test_fixture_station_sheets(fixture_parse: StationsParse) -> None:
    summary = [
        (sheet.sheet_name, sheet.event_date, len(sheet.rows)) for sheet in fixture_parse.sheets
    ]

    assert summary == [
        ("9 6 26", date(2026, 9, 6), 24),
        ("9 13 26", date(2026, 9, 13), 13),
    ]
    assert [sheet.layout for sheet in fixture_parse.sheets] == [FIXTURE_LAYOUT] * 2
    assert fixture_parse.findings == ()


def test_fixture_hadley_total_is_recomputed(fixture_parse: StationsParse) -> None:
    sept_13 = fixture_parse.sheets[1]
    hadley = next(row for row in sept_13.rows if row.raw_name == "Hadley, Ike")

    assert hadley == StationHitsRow(15, "Hadley, Ike", HADLEY_HITS)
    assert sum(hits for _, hits in hadley.hits) == 36


def test_fixture_keeps_raw_names(fixture_parse: StationsParse) -> None:
    sept_6 = fixture_parse.sheets[0]

    assert sept_6.rows[7].row_number == 17
    assert sept_6.rows[7].raw_name == "Kowalczyk, Barrett\xa0"


def test_fixture_matches_golden_snapshot(fixture_parse: StationsParse) -> None:
    snap = json.loads((GOLDEN / "stations_parse.json").read_text())
    # The hash pins the snapshot itself, so a regenerated file cannot slip through.
    canonical = json.dumps(snap, separators=(",", ":"), sort_keys=True).encode()
    assert hashlib.sha256(canonical).hexdigest() == (
        "e892fc5824a14b07d209e20a1f435b36d27699b02f0f3a547b1e281d65c8d2eb"
    )
    got = [
        {
            "sheet": s.sheet_name,
            "event_date": s.event_date.isoformat(),
            "layout": [[e.station_no, e.target_count] for e in s.layout],
            "entries": sorted([name_key(r.raw_name), sum(h for _, h in r.hits)] for r in s.rows),
        }
        for s in sorted(fixture_parse.sheets, key=lambda sheet: sheet.event_date)
    ]
    assert got == snap


def test_stations_fixture_loads_and_parses_without_warnings(
    stations_file_bytes: bytes,
) -> None:
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        parse = parse_stations(load_workbook_bytes(stations_file_bytes))

    assert [str(warning.message) for warning in caught] == []
    assert len(parse.sheets) == 2


def test_uncached_totals_are_recomputed(
    stations_file_bytes: bytes, fixture_parse: StationsParse
) -> None:
    # Re-saving with openpyxl drops every cached SUM/MEDIAN value.
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        workbook = openpyxl.load_workbook(io.BytesIO(stations_file_bytes))
    buffer = io.BytesIO()
    workbook.save(buffer)

    stripped = parse_stations(load_workbook_bytes(buffer.getvalue()))

    assert stripped == fixture_parse


def test_tab_name_a2_mismatch_trusts_a2() -> None:
    parse = _parse({"9 20 26": station_sheet(SUNDAY, [HADLEY])})

    assert [sheet.event_date for sheet in parse.sheets] == [SUNDAY]
    assert _codes(parse) == [("sheet_name_date_mismatch", Severity.WARNING, "9 20 26", None)]
    assert parse.findings[0].event_date == SUNDAY


def test_unreadable_a2_uses_the_tab_name_date() -> None:
    parse = _parse({"9 13 26": station_sheet("no date", [HADLEY])})

    assert [sheet.event_date for sheet in parse.sheets] == [SUNDAY]
    assert _codes(parse) == [("sheet_date_from_name", Severity.WARNING, "9 13 26", None)]


def test_a2_as_text_or_serial_is_read() -> None:
    parse = _parse(
        {
            "9 13 26": station_sheet("9/13/2026", [HADLEY]),
            "9 20 26": station_sheet(46285, [HADLEY]),
        }
    )

    assert [sheet.event_date for sheet in parse.sheets] == [SUNDAY, NEXT_SUNDAY]
    assert parse.findings == ()


def test_sheet_without_any_date_is_left_out() -> None:
    parse = _parse({"Week 3": station_sheet(None, [HADLEY])})

    assert parse.sheets == ()
    assert _codes(parse) == [("sheet_date_missing", Severity.ERROR, "Week 3", None)]


def test_roster_sheet_is_ignored() -> None:
    parse = _parse(
        {
            "9 13 26": station_sheet(SUNDAY, [HADLEY]),
            "Name List (2)": [["Name", "Class", "Other"], ["Hadley, Ike"]],
        }
    )

    assert [sheet.sheet_name for sheet in parse.sheets] == ["9 13 26"]
    assert parse.findings == ()


def test_empty_template_tab_is_left_out_without_displacing_a_week() -> None:
    parse = _parse(
        {
            "9 13 26": station_sheet(SUNDAY, [HADLEY]),
            "9 20 26": station_sheet(NEXT_SUNDAY),
            "9 27 26": station_sheet(SUNDAY),  # copied template, stale A2
        }
    )

    assert [sheet.sheet_name for sheet in parse.sheets] == ["9 13 26"]
    assert _codes(parse) == [
        ("empty_station_sheet", Severity.WARNING, "9 20 26", None),
        ("empty_station_sheet", Severity.WARNING, "9 27 26", None),
    ]


def test_blank_row_mid_table_keeps_later_rows() -> None:
    parse = _parse({"9 13 26": station_sheet(SUNDAY, [HADLEY, [], DEVLIN])})

    assert [row.row_number for row in parse.sheets[0].rows] == [
        FIRST_ENTRY_ROW,
        FIRST_ENTRY_ROW + 2,
    ]
    assert parse.findings == ()


@pytest.mark.parametrize(
    "entry",
    [
        ["Hadley, Ike", 3, 4, None, 5, 7, 4, 7],
        ["Hadley, Ike", 3, 4, "x", 5, 7, 4, 7],
        ["Hadley, Ike", 3, 4, 6, 5, 7, 4, 6.5],
        ["Hadley, Ike", 3, 4, True, 5, 7, 4, 7],  # a boolean is never a number
    ],
)
def test_row_with_some_hits_unreadable_is_dropped(entry: list[object]) -> None:
    parse = _parse({"9 13 26": station_sheet(SUNDAY, [entry, DEVLIN])})

    assert [row.raw_name for row in parse.sheets[0].rows] == ["Devlin, Sid"]
    assert _codes(parse) == [("hits_missing", Severity.ERROR, "9 13 26", 10)]
    assert parse.findings[0].name == "Hadley, Ike"


@pytest.mark.parametrize(
    "entry",
    [
        ["Hadley, Ike", 8, 4, 6, 5, 7, 4, 7],  # 8 hits of 7 targets
        ["Hadley, Ike", 3, 4, 6, 5, 7, 4, -1],
    ],
)
def test_out_of_range_hits_drop_the_row(entry: list[object]) -> None:
    parse = _parse({"9 13 26": station_sheet(SUNDAY, [entry, DEVLIN])})

    assert [row.raw_name for row in parse.sheets[0].rows] == ["Devlin, Sid"]
    assert _codes(parse) == [("hits_out_of_range", Severity.ERROR, "9 13 26", 10)]


def test_layout_summing_to_49_is_kept() -> None:
    parse = _parse({"9 13 26": station_sheet(SUNDAY, [HADLEY], targets=(7, 7, 7, 7, 7, 7, 7))})

    assert sum(entry.target_count for entry in parse.sheets[0].layout) == 49
    assert parse.findings == ()


def test_name_without_hits_is_skipped_with_info() -> None:
    parse = _parse({"9 13 26": station_sheet(SUNDAY, [HADLEY, ["Devlin, Sid"]])})

    assert len(parse.sheets[0].rows) == 1
    assert _codes(parse) == [("station_row_blank", Severity.INFO, "9 13 26", 11)]


def test_hits_without_a_name_are_skipped_with_a_warning() -> None:
    total_only = [None, None, None, None, None, None, None, None, "=SUM(I10:I10)"]
    parse = _parse(
        {
            "9 13 26": station_sheet(
                SUNDAY, [HADLEY, [None, 3, 4, 6, 5, 7, 4, 7], ["**", 1], total_only]
            )
        }
    )

    assert len(parse.sheets[0].rows) == 1
    assert _codes(parse) == [
        ("station_row_missing_name", Severity.WARNING, "9 13 26", 11),
        ("station_row_missing_name", Severity.WARNING, "9 13 26", 12),
    ]


def test_first_name_only_station_row_is_kept_with_info() -> None:
    parse = _parse({"9 13 26": station_sheet(SUNDAY, [["Desmond", 3, 4, 6, 5, 7, 4, 7]])})

    assert [row.raw_name for row in parse.sheets[0].rows] == ["Desmond"]
    assert _codes(parse) == [("first_name_only", Severity.INFO, "9 13 26", 10)]


def test_labels_are_found_by_text_not_position() -> None:
    sheet: list[list[object]] = [
        ["Event Date"],
        [SUNDAY],
        [None, "Station  #", "4", 5.0, 6, "Total", 99],
        [None, "target count", 7, "7", 8],
        [],
        [None, "NAME"],
        [None, "Hadley, Ike", 5, 6, 7, 18],
    ]

    parse = _parse({"9 13 26": sheet})

    assert parse.sheets[0].layout == (
        StationLayoutEntry(4, 7),
        StationLayoutEntry(5, 7),
        StationLayoutEntry(6, 8),
    )
    assert parse.sheets[0].rows == (
        StationHitsRow(7, "Hadley, Ike", (("4", 5), ("5", 6), ("6", 7))),
    )


@pytest.mark.parametrize(
    "broken",
    [
        station_sheet(SUNDAY, [HADLEY], targets=(7, 7, 7, None, 7, 7, 8)),
        station_sheet(SUNDAY, [HADLEY], targets=(7, 7, 7, 0, 7, 7, 8)),
        station_sheet(SUNDAY, [HADLEY], stations=(4, 5, 6, 6, 8, 9, 10)),
        station_sheet(SUNDAY, [HADLEY], stations=("Total",)),
        station_sheet(SUNDAY, [HADLEY], stations=(0, 5, 6, 7, 8, 9, 10)),
        [["Event Date"], [SUNDAY], ["STATION #", 4], ["Name"], HADLEY],
        [["Event Date"], [SUNDAY], ["STATION #", 4], ["TARGET COUNT", 7], HADLEY],
        [["Event Date"], [SUNDAY], ["Name"], HADLEY],
    ],
    ids=[
        "blank-target",
        "zero-target",
        "repeated-station",
        "no-station-numbers",
        "station-zero",
        "no-target-row",
        "no-name-header",
        "no-station-row",
    ],
)
def test_malformed_layout_excludes_only_that_sheet(broken: SheetRows) -> None:
    parse = _parse({"9 6 26": broken, "9 13 26": station_sheet(SUNDAY, [HADLEY])})

    assert [sheet.sheet_name for sheet in parse.sheets] == ["9 13 26"]
    assert _codes(parse) == [("station_layout_missing", Severity.ERROR, "9 6 26", None)]


def test_station_columns_may_run_to_the_last_column() -> None:
    sheet: list[list[object]] = [
        ["Event Date"],
        [SUNDAY],
        ["STATION #", 4, 5],
        ["TARGET COUNT", 7, 7],
        ["Name"],
        ["Hadley, Ike", 3, 4],
    ]

    parse = _parse({"9 13 26": sheet})

    assert parse.sheets[0].rows == (StationHitsRow(6, "Hadley, Ike", (("4", 3), ("5", 4))),)


def test_copied_tab_with_stale_a2_keeps_the_matching_tab() -> None:
    parse = _parse(
        {
            "9 13 26": station_sheet(SUNDAY, [HADLEY]),
            "9 20 26": station_sheet(SUNDAY, [DEVLIN]),
        }
    )

    assert [sheet.sheet_name for sheet in parse.sheets] == ["9 13 26"]
    assert _codes(parse) == [
        ("sheet_name_date_mismatch", Severity.WARNING, "9 20 26", None),
        ("duplicate_sheet_date", Severity.ERROR, "9 20 26", None),
    ]


@pytest.mark.parametrize(
    ("first", "second"),
    [
        ("9 13 26", "9  13 26"),  # both tab names match A2
        ("Sheet1", "Sheet2"),  # neither does
    ],
)
def test_duplicate_sheet_date_leaves_both_out(first: str, second: str) -> None:
    parse = _parse(
        {first: station_sheet(SUNDAY, [HADLEY]), second: station_sheet(SUNDAY, [DEVLIN])}
    )

    assert parse.sheets == ()
    assert [(f.code, f.sheet) for f in parse.findings] == [
        ("duplicate_sheet_date", first),
        ("duplicate_sheet_date", second),
    ]


def test_every_left_out_sheet_has_one_exclusion_finding() -> None:
    parse = _parse(
        {
            "9 6 26": [["Event Date"], [SUNDAY], ["Name"], HADLEY],
            "9 20 26": station_sheet(NEXT_SUNDAY),
            "Week 5": station_sheet(None, [HADLEY]),
            "Sheet1": station_sheet(SUNDAY, [HADLEY]),
            "Sheet2": station_sheet(SUNDAY, [DEVLIN]),
            "10 11 26": station_sheet(date(2026, 10, 11), [DEVLIN]),
        }
    )
    exclusions = [f.sheet or "" for f in parse.findings if f.code in STATION_SHEET_EXCLUSION_CODES]

    assert [sheet.sheet_name for sheet in parse.sheets] == ["10 11 26"]
    assert sorted(exclusions) == sorted(["9 6 26", "9 20 26", "Week 5", "Sheet1", "Sheet2"])
    assert all(f.row is None for f in parse.findings if f.code in STATION_SHEET_EXCLUSION_CODES)


@pytest.mark.parametrize(
    ("column", "value"),
    [("B", " 3 "), ("C", 4.0), ("D", "6")],
    ids=["padded-text", "integral-float", "text"],
)
def test_whole_number_hits_typed_as_text_or_float_are_read(column: str, value: object) -> None:
    lw = loaded_workbook({"9 13 26": station_sheet(SUNDAY, [HADLEY])})
    # Set on the loaded cell: openpyxl writes 4.0 as "4" (%.16g), which reads back as an int,
    # while a file whose XML holds "4.0" loads as a float.
    cell = lw.values["9 13 26"][f"{column}{FIRST_ENTRY_ROW}"]
    cell.value = value
    assert type(cell.value) is type(value)

    parse = parse_stations(lw)

    assert parse.sheets[0].rows == (StationHitsRow(FIRST_ENTRY_ROW, "Hadley, Ike", HADLEY_HITS),)
    assert parse.findings == ()


def _wide_sheet(width: int) -> list[list[object]]:
    return station_sheet(
        SUNDAY,
        [["Hadley, Ike", *[1] * width], ["Devlin, Sid", 1]],  # Devlin would be hits_missing
        stations=range(1, width + 1),
        targets=[2] * width,
    )


def test_thirty_station_columns_are_read() -> None:
    parse = _parse({"9 13 26": _wide_sheet(30)})

    assert len(parse.sheets[0].layout) == 30
    assert parse.sheets[0].rows[0].hits == tuple((str(station), 1) for station in range(1, 31))
    assert _codes(parse) == [("hits_missing", Severity.ERROR, "9 13 26", FIRST_ENTRY_ROW + 1)]


def test_more_than_thirty_station_columns_exclude_the_sheet_before_its_rows() -> None:
    parse = _parse({"9 13 26": _wide_sheet(31)})

    assert parse.sheets == ()
    # Only the layout finding: no row was scanned, so Devlin's row reports nothing.
    assert _codes(parse) == [("station_layout_missing", Severity.ERROR, "9 13 26", None)]
    assert "more than 30 station columns" in parse.findings[0].message


def test_name_column_inside_the_station_columns_excludes_the_sheet() -> None:
    sheet: list[list[object]] = [
        ["Event Date"],
        [SUNDAY],
        ["STATION #", 4, 5],
        ["TARGET COUNT", 7, 7],
        [None, "Name"],  # column B is also station 4
        [None, "Hadley, Ike", 3],
    ]

    parse = _parse({"9 13 26": sheet})

    assert parse.sheets == ()
    assert _codes(parse) == [("station_layout_missing", Severity.ERROR, "9 13 26", None)]
    assert "Name column inside the station columns" in parse.findings[0].message


def test_a_lettered_station_is_its_own_station_sorted_after_its_number() -> None:
    sheet = station_sheet(
        SUNDAY,
        [["Hadley, Ike", 3, 4, 6, 5, 4, 7], ["Devlin, Sid", 2, 4, 6, 5, 3, 6]],
        stations=(4, 5, 6, 7, " 7a ", 8),
        targets=(7, 7, 7, 8, 6, 7),
    )

    parse = _parse({"9 13 26": sheet})

    [kept] = parse.sheets
    assert kept.layout == (
        StationLayoutEntry(4, 7),
        StationLayoutEntry(5, 7),
        StationLayoutEntry(6, 7),
        StationLayoutEntry(7, 8),
        StationLayoutEntry(7, 6, "7A"),
        StationLayoutEntry(8, 7),
    )
    assert [e.label for e in kept.layout] == ["4", "5", "6", "7", "7A", "8"]
    assert kept.rows[0].hits == (("4", 3), ("5", 4), ("6", 6), ("7", 5), ("7A", 4), ("8", 7))
    assert sum(hits for _, hits in kept.rows[0].hits) == 29
    assert parse.findings == ()


def test_the_first_cell_that_is_not_a_station_label_still_ends_the_station_columns() -> None:
    sheet = station_sheet(
        SUNDAY, [["Hadley, Ike", 3, 4, 6]], stations=(4, "7A", 8), targets=(7, 6, 7)
    )
    sheet[3] = ["STATION #", 4, "7A", "Total", 9, "8"]  # nothing after "Total" is read

    parse = _parse({"9 13 26": sheet})

    assert [e.label for e in parse.sheets[0].layout] == ["4", "7A"]


@pytest.mark.parametrize("label", ["7AB", "A7", "7.5", "Total"])
def test_text_that_is_not_a_station_label_ends_the_columns(label: object) -> None:
    sheet = station_sheet(
        SUNDAY, [["Hadley, Ike", 3, 4]], stations=(4, 5, label, 8), targets=(7, 7, 7, 7)
    )

    parse = _parse({"9 13 26": sheet})

    assert [e.label for e in parse.sheets[0].layout] == ["4", "5"]


@pytest.mark.parametrize("number", [0, -3, 100])
def test_a_number_that_cannot_be_a_station_excludes_the_sheet(number: int) -> None:
    sheet = station_sheet(
        SUNDAY, [["Hadley, Ike", 3, 4]], stations=(4, 5, number, 8), targets=(7, 7, 7, 7)
    )

    parse = _parse({"9 13 26": sheet})

    assert parse.sheets == ()
    assert _codes(parse) == [("station_layout_missing", Severity.ERROR, "9 13 26", None)]


def test_the_same_label_in_two_columns_excludes_the_sheet() -> None:
    dup = station_sheet(SUNDAY, [HADLEY], stations=(4, 5, 6, 7, "7A", "7a", 10))
    parse = _parse({"9 6 26": dup, "9 13 26": station_sheet(SUNDAY, [HADLEY])})

    assert [sheet.sheet_name for sheet in parse.sheets] == ["9 13 26"]
    assert _codes(parse) == [("station_layout_missing", Severity.ERROR, "9 6 26", None)]
    assert "same station label in two columns" in parse.findings[0].message


def test_a_hits_finding_names_the_lettered_station() -> None:
    sheet = station_sheet(
        SUNDAY, [["Hadley, Ike", 3, 4, None, 9]], stations=(4, 5, "7A", 8), targets=(7, 7, 6, 7)
    )

    parse = _parse({"9 13 26": sheet})

    [finding] = [f for f in parse.findings if f.code == "hits_missing"]
    assert finding.message.endswith("station 7A")
    sheet = station_sheet(
        SUNDAY, [["Hadley, Ike", 3, 4, 9, 5]], stations=(4, 5, "7A", 8), targets=(7, 7, 6, 7)
    )
    parse = _parse({"9 13 26": sheet})
    [finding] = [f for f in parse.findings if f.code == "hits_out_of_range"]
    assert "station 7A has 9 of 6" in finding.message
