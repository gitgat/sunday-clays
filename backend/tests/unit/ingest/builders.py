"""Synthetic .xlsx workbooks for ingest tests (test utility, not production).

openpyxl writes a str that starts with "=" as a formula with no cached value,
which is exactly the "file written by a script" quirk the parsers handle.
"""

import io
from collections.abc import Mapping, Sequence

from openpyxl import Workbook

from sunday_clays.ingest.workbook import LoadedWorkbook, load_workbook_bytes

SheetRows = Sequence[Sequence[object]]


def workbook_bytes(sheets: Mapping[str, SheetRows]) -> bytes:
    """One sheet per key, in order; each inner sequence is one row from column A."""
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(list(row))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def loaded_workbook(sheets: Mapping[str, SheetRows]) -> LoadedWorkbook:
    return load_workbook_bytes(workbook_bytes(sheets))
