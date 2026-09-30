import hashlib
import io
import json
from collections import Counter
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from sunday_clays.ingest.scores import parse_scores
from sunday_clays.ingest.types import (
    AttendanceRow,
    Finding,
    ParseError,
    ScoreRow,
    ScoresParse,
    Severity,
)
from sunday_clays.ingest.workbook import LoadedWorkbook, load_workbook_bytes

from .builders import loaded_workbook
from .builders_scores import (
    NEXT_SUNDAY,
    SUNDAY,
    scores_sheets,
    scores_workbook,
    scores_workbook_with_empty_text_results,
)

GOLDEN = Path(__file__).resolve().parents[2] / "golden"
VALID_ROW = ["Hadley, Ike", 34, SUNDAY, "Member", None]
EMPTY_TEXT_FORMULA = '=IF(1=1,"",5)'


@pytest.fixture(scope="module")
def fixture_parse(scores_lw: LoadedWorkbook) -> ScoresParse:
    return parse_scores(scores_lw)


def _codes(parse: ScoresParse) -> list[tuple[str, Severity, int | None]]:
    return [(f.code, f.severity, f.row) for f in parse.findings]


def test_fixture_score_rows(fixture_parse: ScoresParse) -> None:
    rows = fixture_parse.score_rows
    dates = {row.event_date for row in rows}

    assert len(rows) == 7480
    assert len(dates) == 311
    assert {day.weekday() for day in dates} == {6}
    assert (min(dates), max(dates)) == (date(2020, 1, 5), date(2026, 9, 27))
    assert rows[0] == ScoreRow(2, "Atherton, Kurt", 40, date(2020, 1, 5), "member", None)
    assert Counter(row.status for row in rows) == {
        "member": 7010,
        "guest": 316,
        "deceased": 154,
    }
    assert Counter(row.gauge_class for row in rows) == {
        None: 7302,
        "12 Gauge": 151,
        "Sub-Gauge": 9,
        "SxS": 9,
        "28 Gauge": 8,  # includes row 4,835's "29 Gauge", read as 28 Gauge
        "20 Gauge": 1,
    }


def test_fixture_keeps_raw_names(fixture_parse: ScoresParse) -> None:
    row = next(r for r in fixture_parse.score_rows if r.row_number == 5601)

    assert row.raw_name == "Kowalczyk, Barrett\xa0"


def test_fixture_parser_findings(fixture_parse: ScoresParse) -> None:
    counts = Counter((f.code, f.severity) for f in fixture_parse.findings)
    corrected = next(f for f in fixture_parse.findings if f.code == "gauge_class_corrected")

    assert counts == {
        ("first_name_only", Severity.INFO): 17,
        ("gauge_class_corrected", Severity.WARNING): 1,
    }
    assert (corrected.row, corrected.name, corrected.event_date) == (
        4835,
        "Eldridge, Tucker",
        date(2024, 9, 1),
    )


def test_fixture_matches_golden_snapshot(fixture_parse: ScoresParse) -> None:
    snap = json.loads((GOLDEN / "scores_by_date.json").read_text())
    # The hash pins the snapshot itself, so a regenerated file cannot slip through.
    canonical = json.dumps(snap, separators=(",", ":"), sort_keys=True).encode()
    assert hashlib.sha256(canonical).hexdigest() == (
        "3f1857aa84e5f58b206e5dafafb9f3fe7f84314a83749a96342eefef631d32cb"
    )
    by_date: dict[str, list[int]] = {}
    for row in fixture_parse.score_rows:
        acc = by_date.setdefault(row.event_date.isoformat(), [0, 0])
        acc[0] += 1
        acc[1] += row.score
    got = [{"event_date": d, "rows": n, "score_sum": s} for d, (n, s) in sorted(by_date.items())]
    assert got == snap


def test_score_cells_as_text_serials_and_floats_are_coerced() -> None:
    parse = parse_scores(
        scores_workbook(
            [
                ["Hadley, Ike", "34", "9/13/2026", " MEMBER ", None],
                ["Devlin, Sid", 25.0, 46278, "Guest", "12 Gauge"],
                ["Nordquist, Sherman", 42, "2026-09-13", None, "  "],
            ]
        )
    )

    assert parse.score_rows == (
        ScoreRow(2, "Hadley, Ike", 34, SUNDAY, "member", None),
        ScoreRow(3, "Devlin, Sid", 25, SUNDAY, "guest", "12 Gauge"),
        ScoreRow(4, "Nordquist, Sherman", 42, SUNDAY, None, None),
    )
    assert parse.findings == ()


def test_blank_rows_are_skipped_silently() -> None:
    parse = parse_scores(
        scores_workbook(
            [
                VALID_ROW,
                [None, None, None, None, None],
                ["  ", "", None, "\xa0", None],
                ["Devlin, Sid", 25, SUNDAY, "Member", None],
                [],
            ]
        )
    )

    assert [row.row_number for row in parse.score_rows] == [2, 5]
    assert parse.findings == ()


@pytest.mark.parametrize(
    ("row", "code"),
    [
        (["Hadley, Ike", None, SUNDAY, "Member", None], "score_missing"),
        (["Hadley, Ike", "forty", SUNDAY, "Member", None], "score_missing"),
        (["Hadley, Ike", 34.5, SUNDAY, "Member", None], "score_missing"),
        (["Hadley, Ike", True, SUNDAY, "Member", None], "score_missing"),
        (["Hadley, Ike", 51, SUNDAY, "Member", None], "score_out_of_range"),
        (["Hadley, Ike", -1, SUNDAY, "Member", None], "score_out_of_range"),
        (["Hadley, Ike", 34, "Sept 13", "Member", None], "unparseable_date"),
        (["Hadley, Ike", 34, None, "Member", None], "unparseable_date"),
        ([None, 34, SUNDAY, "Member", None], "name_missing"),
        (["**", 34, SUNDAY, "Member", None], "name_missing"),
        (["#N/A", 34, SUNDAY, "Member", None], "name_missing"),
        (["#DIV/0!", 34, SUNDAY, "Member", None], "name_missing"),
        # Several defects: only the first in D13 order is reported.
        ([None, "forty", None, "Member", None], "name_missing"),
        (["=B1", None, SUNDAY, "Member", None], "formula_without_cached_value"),
        (
            ["Hadley, Ike", "=30+4", SUNDAY, "Member", None],
            "formula_without_cached_value",
        ),
        (["=B1", 34, SUNDAY, "Member", None], "formula_without_cached_value"),
        (
            ["Hadley, Ike", 34, "=TODAY()", "Member", None],
            "formula_without_cached_value",
        ),
        (
            ["Hadley, Ike", ArrayFormula("B2", "=SUM(30,4)"), SUNDAY, "Member", None],
            "formula_without_cached_value",
        ),
        (
            ["Hadley, Ike", DataTableFormula("B2"), SUNDAY, "Member", None],
            "formula_without_cached_value",
        ),
    ],
)
def test_per_row_errors_drop_only_that_row(row: list[object], code: str) -> None:
    parse = parse_scores(scores_workbook([row, VALID_ROW]))

    assert parse.score_rows == (ScoreRow(3, "Hadley, Ike", 34, SUNDAY, "member", None),)
    assert _codes(parse) == [(code, Severity.ERROR, 2)]
    assert parse.findings[0].sheet == "ALL SCORE DETAIL"


def test_status_outside_the_enum_is_cleared_with_a_warning() -> None:
    parse = parse_scores(scores_workbook([["Hadley, Ike", 34, SUNDAY, "Visitor", None]]))

    assert parse.score_rows[0].status is None
    assert parse.findings == (
        Finding(
            "unknown_status",
            Severity.WARNING,
            "Status 'Visitor' is not Member, Guest or Deceased; it was left blank",
            sheet="ALL SCORE DETAIL",
            row=2,
            event_date=SUNDAY,
            name="Hadley, Ike",
        ),
    )


@pytest.mark.parametrize(
    ("raw", "expected", "code"),
    [
        ("12 Gauge", "12 Gauge", None),
        ("12gauge", "12 Gauge", None),
        (" 20  GAUGE ", "20 Gauge", None),
        ("28 Gauge", "28 Gauge", None),
        (".410", ".410", None),
        ("410", ".410", None),
        ("sub-gauge", "Sub-Gauge", None),
        ("SXS", "SxS", None),
        ("  ", None, None),
        ("29 Gauge", "28 Gauge", "gauge_class_corrected"),
        ("16 Gauge", "16 Gauge", "unknown_gauge_class"),
        (12, "12", "unknown_gauge_class"),
    ],
)
def test_gauge_class_normalization(raw: object, expected: str | None, code: str | None) -> None:
    parse = parse_scores(scores_workbook([["Hadley, Ike", 34, SUNDAY, "Member", raw]]))

    assert parse.score_rows[0].gauge_class == expected
    assert [(f.code, f.severity) for f in parse.findings] == (
        [] if code is None else [(code, Severity.WARNING)]
    )


def test_first_name_only_row_is_kept_with_info() -> None:
    parse = parse_scores(scores_workbook([["Desmond", 30, SUNDAY, "Guest", None]]))

    assert parse.score_rows == (ScoreRow(2, "Desmond", 30, SUNDAY, "guest", None),)
    assert [(f.code, f.severity, f.name, f.event_date) for f in parse.findings] == [
        ("first_name_only", Severity.INFO, "Desmond", SUNDAY)
    ]


def test_header_is_found_by_label_not_position() -> None:
    sheets = scores_sheets(attendance_rows=[[SUNDAY, 1]])
    sheets["ALL SCORE DETAIL"] = [
        ["Sunday Clays - all scores"],
        [],
        ["Event", "class", "  NAME ", "Status", "Score  Shot", "Notes"],
        [SUNDAY, "SxS", "Hadley, Ike", "Member", 34, "windy"],
    ]

    parse = parse_scores(loaded_workbook(sheets))

    assert parse.score_rows == (ScoreRow(4, "Hadley, Ike", 34, SUNDAY, "member", "SxS"),)


def test_missing_header_raises_parse_error() -> None:
    sheets = scores_sheets()
    sheets["ALL SCORE DETAIL"] = [["Name", "Score", "Event"], VALID_ROW]

    with pytest.raises(ParseError) as excinfo:
        parse_scores(loaded_workbook(sheets))

    assert str(excinfo.value) == (
        "The 'ALL SCORE DETAIL' sheet has no header row with the columns "
        "Name, Score Shot, Event, Status, Class"
    )


def test_missing_score_sheet_raises_parse_error() -> None:
    sheets = scores_sheets()
    del sheets["ALL SCORE DETAIL"]

    with pytest.raises(ParseError) as excinfo:
        parse_scores(loaded_workbook(sheets))

    assert str(excinfo.value) == "The workbook has no 'ALL SCORE DETAIL' sheet"


def test_fixture_attendance_rows(fixture_parse: ScoresParse) -> None:
    rows = fixture_parse.attendance_rows

    assert len(rows) == 360
    assert rows[0] == AttendanceRow(2, date(2018, 12, 30), 7)
    assert rows[-1] == AttendanceRow(361, date(2026, 9, 27), 23)


def test_attendance_date_formulas_evaluated(
    scores_file_bytes: bytes, fixture_parse: ScoresParse
) -> None:
    # Re-saving with openpyxl drops every cached formula value, as any script
    # that writes the workbook would.
    workbook = openpyxl.load_workbook(io.BytesIO(scores_file_bytes))
    buffer = io.BytesIO()
    workbook.save(buffer)

    stripped = parse_scores(load_workbook_bytes(buffer.getvalue()))
    codes = Counter(f.code for f in stripped.findings)

    assert len(stripped.attendance_rows) == 360
    assert stripped.attendance_rows == fixture_parse.attendance_rows
    assert codes["date_formula_evaluated"] == 35
    assert codes["formula_without_cached_value"] == 0
    assert stripped.score_rows == fixture_parse.score_rows


def test_uncached_formula_is_error_finding() -> None:
    parse = parse_scores(scores_workbook(attendance_rows=[[SUNDAY, 20], ["=TODAY()", 21]]))

    assert parse.attendance_rows == (AttendanceRow(2, SUNDAY, 20),)
    assert _codes(parse) == [
        ("formula_without_cached_value", Severity.ERROR, 3),
    ]
    assert parse.findings[0].sheet == "Attendance History"


def test_date_formula_chain_passes_through_blank_count_rows() -> None:
    parse = parse_scores(
        scores_workbook(attendance_rows=[[SUNDAY, 20], ["=A2+7", None], ["=$A$3 + 7", 22]])
    )

    assert parse.attendance_rows == (
        AttendanceRow(2, SUNDAY, 20),
        AttendanceRow(4, SUNDAY + timedelta(days=14), 22),
    )
    assert _codes(parse) == [("date_formula_evaluated", Severity.INFO, 4)]


@pytest.mark.parametrize(
    "date_cell",
    [
        "=A2+7",  # row 4 must reference row 3
        "=B3+7",  # not the Date column
        "=A3-7",  # only "+ N days" is evaluated
        "=A3+7*2",  # the whole formula must be "row above + N"
        "=A3+3000000",  # past 9999-12-31
        "=A3+1000000000",  # more days than a timedelta holds
        pytest.param("=A3+" + "9" * 5000, id="5000-digit-days"),  # too long for int()
        pytest.param("=A" + "0" * 4999 + "3+7", id="5000-digit-row"),
    ],
)
def test_other_date_formulas_are_errors(date_cell: str) -> None:
    parse = parse_scores(
        scores_workbook(attendance_rows=[[SUNDAY, 20], [NEXT_SUNDAY, 21], [date_cell, 22]])
    )

    assert [r.row_number for r in parse.attendance_rows] == [2, 3]
    assert _codes(parse) == [("formula_without_cached_value", Severity.ERROR, 4)]


def test_date_formula_needs_a_resolved_row_above() -> None:
    # A1 is the header, so "=A1+7" in row 2 has nothing to add days to.
    parse = parse_scores(scores_workbook(attendance_rows=[["=A1+7", 20]]))

    assert parse.attendance_rows == ()
    assert _codes(parse) == [("formula_without_cached_value", Severity.ERROR, 2)]


def test_uncached_count_formula_is_error_finding() -> None:
    parse = parse_scores(scores_workbook(attendance_rows=[[SUNDAY, "=10+10"]]))

    assert parse.attendance_rows == ()
    assert _codes(parse) == [("formula_without_cached_value", Severity.ERROR, 2)]


def test_attendance_dates_as_text_and_serials_and_blank_counts() -> None:
    parse = parse_scores(
        scores_workbook(
            attendance_rows=[
                ["9/13/2026", 20],
                [46271, " 19 "],
                [date(2026, 9, 27), None],
                [date(2026, 10, 4), "  "],
            ]
        )
    )

    assert parse.attendance_rows == (
        AttendanceRow(2, date(2026, 9, 13), 20),
        AttendanceRow(3, date(2026, 9, 6), 19),
    )
    assert parse.findings == ()


@pytest.mark.parametrize(
    ("row", "code"),
    [
        (["Sept 13", 20], "unparseable_date"),
        ([None, 20], "unparseable_date"),
        ([SUNDAY, "lots"], "head_count_invalid"),
        ([SUNDAY, -3], "head_count_invalid"),
    ],
)
def test_bad_attendance_rows_are_dropped_with_an_error(row: list[object], code: str) -> None:
    parse = parse_scores(scores_workbook(attendance_rows=[row, [NEXT_SUNDAY, 21]]))

    assert parse.attendance_rows == (AttendanceRow(3, NEXT_SUNDAY, 21),)
    assert _codes(parse) == [(code, Severity.ERROR, 2)]


def test_missing_attendance_sheet_raises_parse_error() -> None:
    sheets = scores_sheets([VALID_ROW], [[SUNDAY, 1]])
    del sheets["Attendance History"]

    with pytest.raises(ParseError) as excinfo:
        parse_scores(loaded_workbook(sheets))

    assert str(excinfo.value) == "The workbook has no 'Attendance History' sheet"


def test_formula_cached_as_empty_text_is_a_blank_count() -> None:
    parse = parse_scores(
        scores_workbook_with_empty_text_results(
            attendance_rows=[[SUNDAY, 20], [NEXT_SUNDAY, EMPTY_TEXT_FORMULA]],
            attendance_refs=["B3"],
        )
    )

    assert parse.attendance_rows == (AttendanceRow(2, SUNDAY, 20),)
    assert parse.findings == ()


def test_formula_cached_as_empty_text_is_a_blank_name() -> None:
    parse = parse_scores(
        scores_workbook_with_empty_text_results(
            [VALID_ROW, [EMPTY_TEXT_FORMULA, None, None, None, None]],
            score_refs=["A3"],
        )
    )

    assert parse.score_rows == (ScoreRow(2, "Hadley, Ike", 34, SUNDAY, "member", None),)
    assert parse.findings == ()


@pytest.mark.parametrize(
    ("row", "column"),
    [
        (["Hadley, Ike", 34, SUNDAY, '=IF(1=1,"Member")', None], "Status"),
        (["Hadley, Ike", 34, SUNDAY, "Member", '=IF(1=1,"SxS")'], "Class"),
    ],
)
def test_uncached_status_or_class_formula_is_left_blank_with_a_warning(
    row: list[object], column: str
) -> None:
    parse = parse_scores(scores_workbook([row]))

    kept = parse.score_rows[0]
    assert (kept.row_number, kept.status, kept.gauge_class) == (
        2,
        None if column == "Status" else "member",
        None,
    )
    assert parse.findings == (
        Finding(
            "formula_without_cached_value",
            Severity.WARNING,
            f"{column} holds a formula with no saved value, so it was left blank; "
            "open the file in Excel and save it again to keep it",
            sheet="ALL SCORE DETAIL",
            row=2,
            event_date=SUNDAY,
            name="Hadley, Ike",
        ),
    )


def test_status_and_class_formulas_cached_as_empty_text_are_blank() -> None:
    parse = parse_scores(
        scores_workbook_with_empty_text_results(
            [["Hadley, Ike", 34, SUNDAY, '=IF(1=1,"","Member")', '=IF(1=1,"","SxS")']],
            score_refs=["D2", "E2"],
        )
    )

    assert parse.score_rows == (ScoreRow(2, "Hadley, Ike", 34, SUNDAY, None, None),)
    assert parse.findings == ()


def test_dropped_row_gets_no_status_formula_warning() -> None:
    parse = parse_scores(scores_workbook([[None, 34, SUNDAY, '=IF(1=1,"Member")', None]]))

    assert parse.score_rows == ()
    assert _codes(parse) == [("name_missing", Severity.ERROR, 2)]


def test_excel_error_in_name_is_not_a_shooter() -> None:
    parse = parse_scores(scores_workbook([["#N/A", 34, SUNDAY, "Member", None]]))

    assert parse.score_rows == ()
    assert [(f.code, f.message, f.name) for f in parse.findings] == [
        ("name_missing", "Name holds the Excel error #N/A, not a name", None)
    ]


@pytest.mark.parametrize(
    ("row", "message"),
    [
        (["=B1", 34, SUNDAY, "Member", None], "Name holds a formula"),
        (["=B1", "=30+4", SUNDAY, "Member", None], "Name and Score Shot hold formulas"),
    ],
)
def test_score_formula_errors_name_the_cells(row: list[object], message: str) -> None:
    parse = parse_scores(scores_workbook([row]))

    assert [f.message for f in parse.findings] == [
        f"{message} with no saved value; open the file in Excel and save it again"
    ]


@pytest.mark.parametrize(
    ("row", "message"),
    [
        (["=TODAY()", 21], "Date holds a formula"),
        ([NEXT_SUNDAY, "=10+10"], "Count holds a formula"),
        (["=A2+7", "=10+10"], "Count holds a formula"),  # the date was worked out
        (["=TODAY()", "=10+10"], "Date and Count hold formulas"),
    ],
)
def test_attendance_formula_errors_name_the_cells(row: list[object], message: str) -> None:
    parse = parse_scores(scores_workbook(attendance_rows=[[SUNDAY, 20], row]))

    assert [(f.severity, f.row, f.message) for f in parse.findings] == [
        (
            Severity.ERROR,
            3,
            f"{message} with no saved value; open the file in Excel and save it again",
        )
    ]


def test_missing_attendance_header_raises_parse_error() -> None:
    sheets = scores_sheets([VALID_ROW])
    sheets["Attendance History"] = [["Date", "Head Count", "Median"], [SUNDAY, 20]]

    with pytest.raises(ParseError) as excinfo:
        parse_scores(loaded_workbook(sheets))

    assert str(excinfo.value) == (
        "The 'Attendance History' sheet has no header row with the columns Date, Count"
    )
