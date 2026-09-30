import io
import logging
from datetime import datetime
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session

from sunday_clays.config import Settings, get_settings
from sunday_clays.models import Base

IMPORTS = Base.metadata.tables["imports"]
JOBS = Base.metadata.tables["jobs"]
AUDIT = Base.metadata.tables["audit_log"]
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLSM = "application/vnd.ms-excel.sheet.macroEnabled.12"


def _upload(
    client: TestClient, data: bytes, name: str = "stations_2026-09-27.xlsx"
) -> dict[str, Any]:
    response = client.post("/api/admin/imports", files={"file": (name, data, XLSX)})
    assert response.status_code == 200, response.text
    return response.json()


def _status(session: Session, import_id: int) -> str:
    return str(session.scalar(select(IMPORTS.c.status).where(IMPORTS.c.id == import_id)))


def _actions(session: Session) -> list[str]:
    return list(session.scalars(select(AUDIT.c.action).order_by(AUDIT.c.id)))


def _last_audit(session: Session) -> tuple[str, str, str | None]:
    """(action, role, ip) of the newest audit row; TestClient's peer address is "testclient"."""
    row = session.execute(
        select(AUDIT.c.action, AUDIT.c.role, AUDIT.c.ip).order_by(AUDIT.c.id.desc()).limit(1)
    ).one()
    return row.action, row.role, row.ip


def _set_upload_limit(monkeypatch: pytest.MonkeyPatch, limit: int) -> None:
    monkeypatch.setenv("MAX_UPLOAD_BYTES", str(limit))
    get_settings.cache_clear()


def _without_latest_event(scores_bytes: bytes) -> bytes:
    """The scores fixture minus every ALL SCORE DETAIL row of its newest event (2026-09-27)."""
    workbook = openpyxl.load_workbook(io.BytesIO(scores_bytes))
    sheet = workbook["ALL SCORE DETAIL"]
    event_col = [cell.value for cell in sheet[1]].index("Event") + 1
    dates = [sheet.cell(row, event_col).value for row in range(2, sheet.max_row + 1)]
    newest = max(value for value in dates if isinstance(value, datetime))
    for row in range(sheet.max_row, 1, -1):
        if sheet.cell(row, event_col).value == newest:
            sheet.delete_rows(row)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def test_upload_returns_preview_and_is_audited(
    admin_client: TestClient, stations_bytes: bytes, session: Session
) -> None:
    preview = _upload(admin_client, stations_bytes)
    assert preview["kind"] == "stations"
    assert preview["filename"] == "stations_2026-09-27.xlsx"
    assert preview["duplicate_of"] is None
    assert _status(session, preview["import_id"]) == "pending"
    assert _last_audit(session) == ("imports.stage", "admin", "testclient")


# python-multipart already cuts a drive-letter path (C:\...) to its basename; these reach the route.
@pytest.mark.parametrize(
    "name", ["Users\\bob\\Station Scores.XLSX", "reports/2026/Station Scores.XLSX"]
)
def test_path_filename_is_reduced_to_basename(
    admin_client: TestClient, stations_bytes: bytes, name: str
) -> None:
    preview = _upload(admin_client, stations_bytes, name=name)
    assert preview["filename"] == "Station Scores.XLSX"


def test_nul_in_filename_is_dropped_before_staging(
    admin_client: TestClient, stations_bytes: bytes, session: Session
) -> None:
    # httpx percent-encodes NUL in files= names, so send the multipart body as-is.
    boundary = "sunday-clays-raw-boundary"
    head = (
        f"--{boundary}\r\n"
        'Content-Disposition: form-data; name="file"; filename="a\x00b.xlsx"\r\n'
        f"Content-Type: {XLSX}\r\n\r\n"
    )
    body = head.encode() + stations_bytes + f"\r\n--{boundary}--\r\n".encode()
    response = admin_client.post(
        "/api/admin/imports",
        content=body,
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
    )
    assert response.status_code == 200, response.text
    preview = response.json()
    assert preview["filename"] == "ab.xlsx"
    stored = session.scalar(select(IMPORTS.c.filename).where(IMPORTS.c.id == preview["import_id"]))
    assert stored == "ab.xlsx"
    details = session.scalar(select(AUDIT.c.details).order_by(AUDIT.c.id.desc()).limit(1))
    assert details is not None
    assert details["filename"] == "ab.xlsx"


def test_xlsm_upload_is_accepted_like_xlsx(admin_client: TestClient, stations_bytes: bytes) -> None:
    response = admin_client.post(
        "/api/admin/imports", files={"file": ("stations_2026-09-27.xlsm", stations_bytes, XLSM)}
    )
    assert response.status_code == 200, response.text
    assert response.json()["kind"] == "stations"
    assert response.json()["filename"] == "stations_2026-09-27.xlsm"


def test_long_filename_keeps_its_last_255_characters(
    admin_client: TestClient, stations_bytes: bytes
) -> None:
    preview = _upload(admin_client, stations_bytes, name="s" * 295 + ".xlsx")
    assert len(preview["filename"]) == 255
    assert preview["filename"].endswith(".xlsx")


def test_reupload_reports_duplicate_of(admin_client: TestClient, stations_bytes: bytes) -> None:
    first = _upload(admin_client, stations_bytes)
    second = _upload(admin_client, stations_bytes)
    assert second["duplicate_of"] == first["import_id"]


def test_get_import_returns_the_staged_preview(
    admin_client: TestClient, stations_bytes: bytes
) -> None:
    preview = _upload(admin_client, stations_bytes)
    response = admin_client.get(f"/api/admin/imports/{preview['import_id']}")
    assert response.status_code == 200
    assert response.json() == preview


def test_list_imports_is_newest_first(
    admin_client: TestClient, stations_bytes: bytes, scores_bytes: bytes
) -> None:
    older = _upload(admin_client, stations_bytes)
    newer = _upload(admin_client, scores_bytes, name="scores_2026-09-27.xlsx")
    listed = admin_client.get("/api/admin/imports").json()
    assert [row["id"] for row in listed] == [newer["import_id"], older["import_id"]]
    assert {row["status"] for row in listed} == {"pending"}
    assert len(listed[0]["sha256"]) == 64


@pytest.mark.parametrize("name", ["scores.csv", "scores.xls", "scores"])
def test_non_xlsx_filename_is_rejected(
    admin_client: TestClient, stations_bytes: bytes, name: str, session: Session
) -> None:
    response = admin_client.post("/api/admin/imports", files={"file": (name, stations_bytes, XLSX)})
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "unsupported_file_type"
    assert session.execute(select(IMPORTS.c.id)).all() == []


def test_unreadable_workbook_is_400_and_stages_nothing(
    admin_client: TestClient, session: Session, caplog: pytest.LogCaptureFixture
) -> None:
    route_logger = "sunday_clays.api.routes.admin_imports"
    with caplog.at_level(logging.WARNING, logger=route_logger):
        response = admin_client.post(
            "/api/admin/imports", files={"file": ("junk.xlsx", b"this is not a zip file", XLSX)}
        )
    assert response.status_code == 400
    assert response.json()["error"] == {
        "code": "unreadable_file",
        "message": "This file could not be read as an Excel workbook",
    }
    # The chained cause (the zipfile error) is logged server-side, never sent to the client.
    records = [r for r in caplog.records if r.name == route_logger]
    assert len(records) == 1
    exc_info = records[0].exc_info
    assert exc_info is not None
    _, exc, _ = exc_info
    assert exc is not None
    assert exc.__cause__ is not None
    assert session.execute(select(IMPORTS.c.id)).all() == []
    assert "imports.stage" not in _actions(session)


def test_upload_backstop_rejects_one_byte_over_limit(
    admin_client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    _set_upload_limit(monkeypatch, 2048)
    over = admin_client.post("/api/admin/imports", files={"file": ("a.xlsx", b"x" * 2049, XLSX)})
    at_limit = admin_client.post(
        "/api/admin/imports", files={"file": ("a.xlsx", b"x" * 2048, XLSX)}
    )
    assert over.status_code == 413
    assert over.json()["error"]["code"] == "payload_too_large"
    assert at_limit.status_code == 400  # size accepted; the parser rejects the bytes
    assert at_limit.json()["error"]["code"] == "unreadable_file"


def test_oversized_upload_without_session_gets_413(
    anon_client: TestClient, auth_env: Settings
) -> None:
    body = b"x" * (auth_env.max_upload_bytes + 65_536 + 1)
    response = anon_client.post("/api/admin/imports", files={"file": ("big.xlsx", body, XLSX)})
    assert response.status_code == 413
    assert response.json()["error"]["code"] == "payload_too_large"


def test_commit_enqueues_rebuild_and_is_audited(
    admin_client: TestClient, stations_bytes: bytes, session: Session
) -> None:
    import_id = _upload(admin_client, stations_bytes)["import_id"]
    response = admin_client.post(
        f"/api/admin/imports/{import_id}/commit", json={"confirm_removals": False}
    )
    assert response.status_code == 200
    job_id = response.json()["job_id"]
    job = session.execute(
        select(JOBS.c.kind, JOBS.c.dedupe_key, JOBS.c.status).where(JOBS.c.id == job_id)
    ).one()
    assert tuple(job) == ("rebuild", "rebuild", "queued")
    assert _status(session, import_id) == "committed"
    assert _last_audit(session) == ("imports.commit", "admin", "testclient")


def test_commit_without_body_does_not_confirm_removals(
    admin_client: TestClient, stations_bytes: bytes, session: Session
) -> None:
    import_id = _upload(admin_client, stations_bytes)["import_id"]
    assert admin_client.post(f"/api/admin/imports/{import_id}/commit").status_code == 200
    assert _status(session, import_id) == "committed"


def test_stale_scores_upload_needs_confirmation_to_commit(
    fx_admin_client: TestClient, scores_bytes: bytes
) -> None:
    preview = _upload(fx_admin_client, _without_latest_event(scores_bytes), name="scores_old.xlsx")
    assert preview["kind"] == "scores"
    assert preview["requires_removal_confirmation"] is True
    commit_url = f"/api/admin/imports/{preview['import_id']}/commit"
    refused = fx_admin_client.post(commit_url, json={"confirm_removals": False})
    assert refused.status_code == 409
    assert refused.json()["error"]["code"] == "removals_not_confirmed"
    assert fx_admin_client.post(commit_url, json={"confirm_removals": True}).status_code == 200


def test_second_commit_is_409_not_pending(admin_client: TestClient, stations_bytes: bytes) -> None:
    import_id = _upload(admin_client, stations_bytes)["import_id"]
    admin_client.post(f"/api/admin/imports/{import_id}/commit")
    again = admin_client.post(f"/api/admin/imports/{import_id}/commit")
    assert again.status_code == 409
    assert again.json()["error"]["code"] == "not_pending"


def test_discard_marks_import_discarded(
    admin_client: TestClient, stations_bytes: bytes, session: Session
) -> None:
    import_id = _upload(admin_client, stations_bytes)["import_id"]
    response = admin_client.post(f"/api/admin/imports/{import_id}/discard")
    assert response.status_code == 204
    assert _status(session, import_id) == "discarded"
    assert _last_audit(session) == ("imports.discard", "admin", "testclient")


def test_discard_of_committed_import_is_409_not_pending(
    admin_client: TestClient, stations_bytes: bytes, session: Session
) -> None:
    import_id = _upload(admin_client, stations_bytes)["import_id"]
    assert admin_client.post(f"/api/admin/imports/{import_id}/commit").status_code == 200
    response = admin_client.post(f"/api/admin/imports/{import_id}/discard")
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "not_pending"
    assert _status(session, import_id) == "committed"
    assert "imports.discard" not in _actions(session)


def test_rollback_marks_import_rolled_back_and_enqueues_rebuild(
    admin_client: TestClient, stations_bytes: bytes, session: Session
) -> None:
    import_id = _upload(admin_client, stations_bytes)["import_id"]
    commit_job_id = admin_client.post(f"/api/admin/imports/{import_id}/commit").json()["job_id"]
    response = admin_client.post(f"/api/admin/imports/{import_id}/rollback")
    assert response.status_code == 200
    # No worker runs here, so commit's rebuild is still queued and the rollback coalesces into it.
    assert response.json()["job_id"] == commit_job_id
    jobs = session.execute(select(JOBS.c.id, JOBS.c.kind, JOBS.c.dedupe_key, JOBS.c.status)).all()
    assert [tuple(job) for job in jobs] == [(commit_job_id, "rebuild", "rebuild", "queued")]
    assert _status(session, import_id) == "rolled_back"
    assert _last_audit(session) == ("imports.rollback", "admin", "testclient")


def test_rollback_of_pending_import_is_409_not_committed(
    admin_client: TestClient, stations_bytes: bytes, session: Session
) -> None:
    import_id = _upload(admin_client, stations_bytes)["import_id"]
    response = admin_client.post(f"/api/admin/imports/{import_id}/rollback")
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "not_committed"
    assert _status(session, import_id) == "pending"
    assert session.execute(select(JOBS.c.id)).all() == []
    assert "imports.rollback" not in _actions(session)


@pytest.mark.parametrize(
    ("method", "path"),
    [
        ("GET", "/api/admin/imports/987654"),
        ("POST", "/api/admin/imports/987654/commit"),
        ("POST", "/api/admin/imports/987654/discard"),
        ("POST", "/api/admin/imports/987654/rollback"),
    ],
)
def test_unknown_import_is_404(admin_client: TestClient, method: str, path: str) -> None:
    response = admin_client.request(method, path)
    assert response.status_code == 404
    assert response.json()["error"]["code"] == "import_not_found"


def test_job_status_reports_the_queued_job(admin_client: TestClient, stations_bytes: bytes) -> None:
    import_id = _upload(admin_client, stations_bytes)["import_id"]
    job_id = admin_client.post(f"/api/admin/imports/{import_id}/commit").json()["job_id"]
    job = admin_client.get(f"/api/admin/jobs/{job_id}").json()
    assert (job["id"], job["kind"], job["status"]) == (job_id, "rebuild", "queued")


def test_unknown_job_is_404(admin_client: TestClient) -> None:
    response = admin_client.get("/api/admin/jobs/987654")
    assert response.status_code == 404
    assert response.json()["error"]["code"] == "job_not_found"
