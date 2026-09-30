"""Domain test helpers: synthetic workbooks in the real fixture layout, commit and seed shortcuts.

Helpers are exposed as fixtures because pytest runs with ``--import-mode=importlib``, where test
modules cannot import each other.
"""

import hashlib
import io
import warnings
from collections.abc import Callable, Iterator, Sequence
from contextlib import AbstractContextManager, contextmanager
from datetime import date, datetime
from typing import Any

import openpyxl
import pytest
from sqlalchemy import event, func, select
from sqlalchemy.orm import Session

from sunday_clays.domain.identity import lookup_shooter, resolve_shooter
from sunday_clays.domain.schemas import StationsDiff
from sunday_clays.ingest.names import identity_key, name_key
from sunday_clays.models import (
    Base,
    Event,
    Import,
    ImportStationHit,
    ImportStationSheet,
    Round,
    StationHit,
)

ScoreSpec = tuple[object, ...]  # (raw_name, score, event_date[, status[, gauge_class]])
HitSpec = tuple[str, Sequence[int]]  # (raw_name, hits per station in layout order)
FIRST_STATION = 4
LiveRows = dict[str, list[dict[str, Any]]]
# every live table (C4), listed here rather than read from rebuild.LIVE_TABLES
_LIVE = (
    "events",
    "rounds",
    "shooter_profiles",
    "station_layouts",
    "station_hits",
    "data_issues",
    "round_metrics",
)
_SURROGATE_IDS = ("rounds", "station_hits", "data_issues")


def _scores_workbook(
    rows: Sequence[ScoreSpec], attendance: Sequence[tuple[date, int]] = ()
) -> bytes:
    """Scores workbook in the fixture layout: ALL SCORE DETAIL + Attendance History."""
    wb = openpyxl.Workbook()
    detail = wb.active
    assert detail is not None
    detail.title = "ALL SCORE DETAIL"
    detail.append(["Name", "Score Shot", "Event", "Status", "Class"])
    for spec in rows:
        status = spec[3] if len(spec) > 3 else "Member"
        gauge = spec[4] if len(spec) > 4 else None
        detail.append([spec[0], spec[1], spec[2], status, gauge])
    history = wb.create_sheet("Attendance History")
    history.append(["Date", "Count", "Median"])
    for event_date, count in attendance:
        history.append([event_date, count, None])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _stations_workbook(
    sheets: Sequence[tuple[str, date, Sequence[int], Sequence[HitSpec]]],
    stations: Sequence[int] = (4, 5, 6, 7, 8, 9, 10),
) -> bytes:
    """Stations workbook with one tab per (tab name, A2 date, target counts, [(name, hits)])."""
    wb = openpyxl.Workbook()
    default = wb.active
    assert default is not None
    wb.remove(default)
    for tab, event_date, targets, entries in sheets:
        ws = wb.create_sheet(tab)
        ws["A1"] = "Event Date"
        ws["A2"] = event_date
        ws.append([])
        ws.append(["STATION #", *stations, "Total"])
        ws.append(["TARGET COUNT", *targets, sum(targets)])
        ws.append([])
        ws.append(["Median Hits"])
        ws.append([])
        ws.append(["Name", "Station Hits", *([None] * (len(stations) - 1)), " Total Hits"])
        for rank, (raw_name, hits) in enumerate(entries, start=1):
            ws.append([raw_name, *hits, sum(hits), rank])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _mark_committed(session: Session, import_id: int) -> None:
    """Flip a staged import to committed (Task 6's commit_import does this for real)."""
    imp = session.get(Import, import_id)
    assert imp is not None
    imp.status = "committed"
    imp.committed_at = func.clock_timestamp()
    session.flush()


def _seed_live_round(
    session: Session, raw_name: str, event_date: date, score: int, ordinal: int = 1
) -> int:
    """Insert one live round (and its event) directly, as a rebuild would have."""
    if session.get(Event, event_date) is None:
        session.add(
            Event(
                event_date=event_date,
                round_type="sporting",
                round_type_source="none",
                n_rounds=1,
                n_shooters=1,
                has_scores=True,
                has_stations=False,
                results_complete=True,
            )
        )
        session.flush()
    rnd = Round(
        event_date=event_date,
        shooter_id=resolve_shooter(session, raw_name, event_date),
        name_key=identity_key(name_key(raw_name), event_date),
        ordinal=ordinal,
        score=score,
        source_row=100 + ordinal,
    )
    session.add(rnd)
    session.flush()
    return rnd.id


def _seed_live_station_entry(
    session: Session, raw_name: str, event_date: date, hits: Sequence[int], entry_row: int
) -> None:
    """Insert a committed station sheet row plus its live station_hits, as a rebuild would.

    Entries on one date share one committed seed import and sheet (tab "seed"), whose stored
    preview is a valid StationsDiff.
    """
    sha256 = hashlib.sha256(f"seed {event_date.isoformat()}".encode()).hexdigest()
    sheet = session.scalar(
        select(ImportStationSheet)
        .join(Import, Import.id == ImportStationSheet.import_id)
        .where(Import.sha256 == sha256)
    )
    if sheet is None:
        diff = StationsDiff(
            events_added=[event_date], events_replaced=[], events_unchanged=[], sheets_skipped=[]
        )
        imp = Import(
            kind="stations",
            filename="seed.xlsx",
            sha256=sha256,
            file_bytes=b"seed",
            status="committed",
            summary={"diff": diff.model_dump(mode="json"), "requires_removal_confirmation": False},
        )
        session.add(imp)
        session.flush()
        sheet = ImportStationSheet(import_id=imp.id, sheet_name="seed", event_date=event_date)
        session.add(sheet)
        session.flush()
    key = identity_key(name_key(raw_name), event_date)
    for station_no, value in enumerate(hits, start=FIRST_STATION):
        session.add(
            ImportStationHit(
                sheet_id=sheet.id,
                row_number=entry_row,
                raw_name=raw_name,
                name_key=key,
                station_no=station_no,
                hits=value,
            )
        )
        session.add(
            StationHit(
                event_date=event_date,
                station_no=station_no,
                sheet_id=sheet.id,
                entry_row=entry_row,
                name_key=key,
                shooter_id=lookup_shooter(session, key),
                hits=value,
            )
        )
    session.flush()


@contextmanager
def _count_queries(session: Session) -> Iterator[list[str]]:
    """Record every SQL statement the session's connection sends inside the block."""
    statements: list[str] = []

    def record(_conn: object, _cursor: object, statement: str, *_rest: object) -> None:
        statements.append(statement)

    connection = session.connection()
    event.listen(connection, "before_cursor_execute", record)
    try:
        yield statements
    finally:
        event.remove(connection, "before_cursor_execute", record)


@pytest.fixture
def scores_workbook() -> Callable[..., bytes]:
    return _scores_workbook


@pytest.fixture
def stations_workbook() -> Callable[..., bytes]:
    return _stations_workbook


@pytest.fixture
def mark_committed() -> Callable[[Session, int], None]:
    return _mark_committed


@pytest.fixture
def seed_live_round() -> Callable[..., int]:
    return _seed_live_round


@pytest.fixture
def seed_live_station_entry() -> Callable[..., None]:
    return _seed_live_station_entry


def _live_rows(session: Session) -> LiveRows:
    """Every row of every live table, surrogate ids renumbered 1..n in id (insert) order.

    Rebuilds never reuse ids, so each round reference (station_hits.round_id,
    round_metrics.round_id, data_issues.details.round_id) is renumbered with its round: two
    rebuilds of the same inputs compare equal, and a shifted station link still shows.
    """
    rows: LiveRows = {}
    for name in _LIVE:
        table = Base.metadata.tables[name]
        found = session.execute(select(table).order_by(*table.primary_key))
        rows[name] = [dict(row._mapping) for row in found]
    round_no = {row["id"]: n for n, row in enumerate(rows["rounds"], start=1)}
    for name in _SURROGATE_IDS:
        for n, row in enumerate(rows[name], start=1):
            row["id"] = n
    for row in (*rows["station_hits"], *rows["round_metrics"]):
        if row["round_id"] is not None:
            row["round_id"] = round_no[row["round_id"]]
    for row in rows["data_issues"]:
        if "round_id" in row["details"]:
            row["details"] = row["details"] | {"round_id": round_no[row["details"]["round_id"]]}
    return rows


@pytest.fixture
def count_queries() -> Callable[[Session], AbstractContextManager[list[str]]]:
    return _count_queries


@pytest.fixture
def live_rows() -> Callable[[Session], LiveRows]:
    return _live_rows


def _edit_workbook(data: bytes, edit: Callable[[Any], None]) -> bytes:
    """Load with cached values (formulas become plain values), apply ``edit``, save."""
    with warnings.catch_warnings():
        # Only openpyxl's data-validation notice (as ingest/workbook.py); any other warning stays
        # visible to the suite's filterwarnings=error.
        warnings.filterwarnings(
            "ignore",
            message="Data Validation extension is not supported",
            category=UserWarning,
            module="openpyxl",
        )
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    edit(wb)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _scores_without_dates(data: bytes, dates: Sequence[date]) -> bytes:
    """Blank every ALL SCORE DETAIL row on ``dates`` (an older copy missing those weeks)."""

    def edit(wb: Any) -> None:
        for row in wb["ALL SCORE DETAIL"].iter_rows(min_row=2):
            event = row[2].value
            if isinstance(event, datetime) and event.date() in dates:
                for cell in row[:5]:
                    cell.value = None

    return _edit_workbook(data, edit)


def _scores_with_score(data: bytes, raw_name: str, event_date: date, score: int) -> bytes:
    """Change the Score Shot of the single row for ``raw_name`` on ``event_date``."""

    def edit(wb: Any) -> None:
        [row] = [
            r
            for r in wb["ALL SCORE DETAIL"].iter_rows(min_row=2)
            if r[0].value == raw_name
            and isinstance(r[2].value, datetime)
            and r[2].value.date() == event_date
        ]
        row[1].value = score

    return _edit_workbook(data, edit)


def _stations_edited(
    data: bytes, drop_tabs: Sequence[str] = (), hit: tuple[str, str, int, int] | None = None
) -> bytes:
    """Drop tabs and/or set one hit; hit = (tab, raw_name, column (1 = first station), value)."""

    def edit(wb: Any) -> None:
        for tab in drop_tabs:
            del wb[tab]
        if hit is not None:
            tab, raw_name, column, value = hit
            [row] = [r for r in wb[tab].iter_rows(min_row=10) if r[0].value == raw_name]
            row[column].value = value

    return _edit_workbook(data, edit)


@pytest.fixture
def scores_without_dates() -> Callable[..., bytes]:
    return _scores_without_dates


@pytest.fixture
def scores_with_score() -> Callable[..., bytes]:
    return _scores_with_score


@pytest.fixture
def stations_edited() -> Callable[..., bytes]:
    return _stations_edited
