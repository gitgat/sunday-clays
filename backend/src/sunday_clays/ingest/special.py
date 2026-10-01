"""Special-event workbooks (Plan 17): one sheet, one Sunday, one round per shooter.

The "Special Event" sheet (titles match ignoring case and spacing):

    A1 "Special event"  B1 the shoot's name, e.g. "Three Clay Shoot"
    A2 "Event date"     B2 the date
    A3 "Station"        B3... station labels (1, 2, ... or 7A), then optionally "Total"
    A4 "Targets"        B4... targets thrown at each station
    A5...               one shooter per row: name, hits per station, optionally the total

Every total is recomputed from the station hits; a total cell that disagrees is reported.
Pure module: no database, no I/O.
"""

from __future__ import annotations

from datetime import date

from openpyxl.worksheet.worksheet import Worksheet

from sunday_clays.ingest.names import clean_display_name
from sunday_clays.ingest.types import (
    Finding,
    ParseError,
    Severity,
    SpecialParse,
    SpecialRow,
    StationLayoutEntry,
)
from sunday_clays.ingest.workbook import (
    SPECIAL_SHEET,
    LoadedWorkbook,
    coerce_date,
    coerce_int,
    find_sheet,
    is_blank,
    name_text,
    normalize_label,
)
from sunday_clays.station_label import label_number, parse_label

LABEL_ROW, DATE_ROW, STATION_ROW, TARGET_ROW, FIRST_ENTRY_ROW = 1, 2, 3, 4, 5
VALUE_COLUMN = 2  # B: the label, the date and the first station
MAX_STATIONS = 30
MAX_LABEL_CHARS = 60
TOTAL_HEADER = "total"

MISSING_SHEET = f'The workbook has no "{SPECIAL_SHEET}" sheet'
MISSING_LABEL = 'Cell B1 needs the name of the special shoot, such as "Three Clay Shoot"'
LABEL_TOO_LONG = f"The special shoot's name in B1 is longer than {MAX_LABEL_CHARS} characters"
MISSING_DATE = "Cell B2 needs the date of the special shoot"
NO_STATIONS = "Row 3 needs station numbers from column B (1, 2, 3 or a label such as 7A)"
BAD_TARGETS = "Row 4 needs a target count of 1 or more under every station"
TOO_MANY_STATIONS = f"Row 3 has more than {MAX_STATIONS} stations"
NO_ROWS = "No shooter rows could be read from row 5 down"


def parse_special(lw: LoadedWorkbook) -> SpecialParse:
    """The special sheet as a SpecialParse; ParseError for anything that cannot be used."""
    sheet = find_sheet(lw.values, SPECIAL_SHEET)
    if sheet is None:
        raise ParseError(MISSING_SHEET)
    label = _label(sheet)
    event_date = coerce_date(sheet.cell(row=DATE_ROW, column=VALUE_COLUMN).value)
    if event_date is None:
        raise ParseError(MISSING_DATE)
    layout, total_column = _layout(sheet)
    findings: list[Finding] = []
    rows = [
        row
        for number in range(FIRST_ENTRY_ROW, sheet.max_row + 1)
        if (row := _row(sheet, number, layout, total_column, event_date, findings)) is not None
    ]
    if not rows:
        if findings:
            first = findings[0]
            raise ParseError(f"{NO_ROWS}: {first.message} (row {first.row})")
        raise ParseError(NO_ROWS)
    return SpecialParse(event_date, label, layout, tuple(rows), tuple(findings))


def _label(sheet: Worksheet) -> str:
    value = sheet.cell(row=LABEL_ROW, column=VALUE_COLUMN).value
    text = "" if is_blank(value) else clean_display_name(str(value))
    if not text:
        raise ParseError(MISSING_LABEL)
    if len(text) > MAX_LABEL_CHARS:
        raise ParseError(LABEL_TOO_LONG)
    return text


def _layout(sheet: Worksheet) -> tuple[tuple[StationLayoutEntry, ...], int | None]:
    """Station columns from B3 rightwards, up to a blank or a "Total" header (its column)."""
    entries: list[StationLayoutEntry] = []
    for column in range(VALUE_COLUMN, VALUE_COLUMN + MAX_STATIONS + 1):
        header = sheet.cell(row=STATION_ROW, column=column).value
        if is_blank(header):
            break
        if normalize_label(header) == TOTAL_HEADER:
            return _checked(entries), column
        label = parse_label(header)
        if label is None:
            raise ParseError(f"Row 3 has {header!r} where a station number belongs")
        if any(entry.label == label for entry in entries):
            raise ParseError(f"Station {label} appears twice in row 3")
        target = coerce_int(sheet.cell(row=TARGET_ROW, column=column).value)
        if target is None or target < 1:
            raise ParseError(BAD_TARGETS)
        entries.append(StationLayoutEntry(label_number(label), target, label))
    return _checked(entries), None


def _checked(entries: list[StationLayoutEntry]) -> tuple[StationLayoutEntry, ...]:
    if not entries:
        raise ParseError(NO_STATIONS)
    if len(entries) > MAX_STATIONS:
        raise ParseError(TOO_MANY_STATIONS)
    return tuple(entries)


def _shown(value: object) -> str:
    return "nothing" if is_blank(value) else str(value)


def _row(
    sheet: Worksheet,
    number: int,
    layout: tuple[StationLayoutEntry, ...],
    total_column: int | None,
    event_date: date,
    findings: list[Finding],
) -> SpecialRow | None:
    raw_name = name_text(sheet.cell(row=number, column=1).value)
    values = [sheet.cell(row=number, column=VALUE_COLUMN + i).value for i in range(len(layout))]
    if raw_name is None:
        if not all(is_blank(value) for value in values):
            findings.append(
                Finding(
                    "special_row_without_name",
                    Severity.WARNING,
                    "This row has hits but no name; it is left out",
                    sheet=SPECIAL_SHEET,
                    row=number,
                    event_date=event_date,
                )
            )
        return None
    name = clean_display_name(raw_name)
    hits: list[tuple[str, int]] = []
    for entry, value in zip(layout, values, strict=True):
        count = coerce_int(value)
        if count is None or not 0 <= count <= entry.target_count:
            findings.append(
                Finding(
                    "special_hits_invalid",
                    Severity.ERROR,
                    f"Station {entry.label} has {_shown(value)}, not a hit count from 0 to "
                    f"{entry.target_count}; the row is left out",
                    sheet=SPECIAL_SHEET,
                    row=number,
                    event_date=event_date,
                    name=name,
                )
            )
            return None
        hits.append((entry.label, count))
    row = SpecialRow(number, raw_name, tuple(hits))
    if total_column is not None:
        given = sheet.cell(row=number, column=total_column).value
        if not is_blank(given) and coerce_int(given) != row.total:
            findings.append(
                Finding(
                    "special_total_mismatch",
                    Severity.WARNING,
                    f"The total column says {_shown(given)} but the stations add up to "
                    f"{row.total}; {row.total} is used",
                    sheet=SPECIAL_SHEET,
                    row=number,
                    event_date=event_date,
                    name=name,
                )
            )
    return row
