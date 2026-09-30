"""Parse the cumulative scores & attendance workbook (Contract C3).

Only the parser drops data, and every row it drops gets an ERROR finding.
"""

import re
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date, timedelta

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from openpyxl.worksheet.worksheet import Worksheet

from sunday_clays.ingest.names import clean_display_name, name_key
from sunday_clays.ingest.types import (
    AttendanceRow,
    Finding,
    ParseError,
    ScoreRow,
    ScoresParse,
    Severity,
)
from sunday_clays.ingest.workbook import (
    SCORES_SHEET,
    LoadedWorkbook,
    coerce_date,
    coerce_int,
    find_sheet,
    is_blank,
    name_text,
    normalize_label,
)

ATTENDANCE_SHEET = "Attendance History"
SCORE_HEADER = ("Name", "Score Shot", "Event", "Status", "Class")
ATTENDANCE_HEADER = ("Date", "Count")
HEADER_SEARCH_ROWS = 20
HEADER_SEARCH_COLUMNS = 30
MAX_SCORE = 50
RESAVE_HINT = "open the file in Excel and save it again"

_REQUIRED_SCORE_CELLS = ("Name", "Score Shot", "Event")
_STATUSES = frozenset({"member", "guest", "deceased"})
_GAUGES = {
    "12gauge": "12 Gauge",
    "20gauge": "20 Gauge",
    "28gauge": "28 Gauge",
    ".410": ".410",
    "410": ".410",
    "sub-gauge": "Sub-Gauge",
    "sxs": "SxS",
}
_CORRECTED_GAUGES = {"29gauge": "28 Gauge"}  # user decision 2026-09-28: a 28 ga typo

MakeFinding = Callable[[str, Severity, str], Finding]


@dataclass(frozen=True)
class _Cell:
    value: object  # cached value (data_only=True)
    formula: object  # formula text, or the same value when there is no formula
    data_type: str  # openpyxl type of the cached value ("e" = an Excel error)

    @property
    def uncached(self) -> bool:
        # Excel and LibreOffice save a formula that returned "" as t="str" with
        # an empty <v>, which openpyxl reads as None: a cached blank, not a
        # missing value. openpyxl-written formulas carry no type and read as "n".
        return self.value is None and self.data_type != "str" and _is_formula(self.formula)

    @property
    def error(self) -> bool:
        return self.data_type == "e"

    @property
    def blank(self) -> bool:
        return is_blank(self.value) and not self.uncached


def parse_scores(lw: LoadedWorkbook) -> ScoresParse:
    findings: list[Finding] = []
    score_rows = _parse_score_sheet(lw, findings)
    attendance_rows = _parse_attendance_sheet(lw, findings)
    return ScoresParse(
        score_rows=tuple(score_rows),
        attendance_rows=tuple(attendance_rows),
        findings=tuple(findings),
    )


def _is_formula(value: object) -> bool:
    if isinstance(value, ArrayFormula | DataTableFormula):
        return True
    return isinstance(value, str) and value.startswith("=")


def _uncached_message(labels: list[str]) -> str:
    held = "holds a formula" if len(labels) == 1 else "hold formulas"
    return f"{' and '.join(labels)} {held} with no saved value; {RESAVE_HINT}"


def _sheet_pair(lw: LoadedWorkbook, title: str) -> tuple[Worksheet, Worksheet]:
    values = find_sheet(lw.values, title)
    formulas = find_sheet(lw.formulas, title)
    if values is None or formulas is None:
        raise ParseError(f"The workbook has no '{title}' sheet")
    return values, formulas


def _find_header(sheet: Worksheet, labels: tuple[str, ...]) -> tuple[int, dict[str, int]]:
    wanted = {normalize_label(label): label for label in labels}
    last_row = min(sheet.max_row, HEADER_SEARCH_ROWS)
    last_column = min(sheet.max_column, HEADER_SEARCH_COLUMNS)
    for row in range(1, last_row + 1):
        columns: dict[str, int] = {}
        for column in range(1, last_column + 1):
            label = wanted.get(normalize_label(sheet.cell(row, column).value))
            if label is not None and label not in columns:
                columns[label] = column
        if len(columns) == len(labels):
            return row, columns
    raise ParseError(
        f"The '{sheet.title}' sheet has no header row with the columns " + ", ".join(labels)
    )


def _cell(values: Worksheet, formulas: Worksheet, row: int, column: int) -> _Cell:
    cached = values.cell(row, column)
    return _Cell(cached.value, formulas.cell(row, column).value, cached.data_type)


def _parse_score_sheet(lw: LoadedWorkbook, findings: list[Finding]) -> list[ScoreRow]:
    values, formulas = _sheet_pair(lw, SCORES_SHEET)
    header_row, columns = _find_header(values, SCORE_HEADER)
    rows: list[ScoreRow] = []
    for row in range(header_row + 1, values.max_row + 1):
        cells = {label: _cell(values, formulas, row, column) for label, column in columns.items()}
        if all(cell.blank for cell in cells.values()):
            continue
        score_row = _score_row(row, cells, findings)
        if score_row is not None:
            rows.append(score_row)
    return rows


def _score_row(row: int, cells: dict[str, _Cell], findings: list[Finding]) -> ScoreRow | None:
    name_cell = cells["Name"]
    raw_name = None if name_cell.error else name_text(name_cell.value)
    event_date = coerce_date(cells["Event"].value)

    def finding(code: str, severity: Severity, message: str) -> Finding:
        return Finding(
            code,
            severity,
            message,
            sheet=SCORES_SHEET,
            row=row,
            event_date=event_date,
            name=clean_display_name(raw_name) if raw_name is not None else None,
        )

    uncached = [label for label in _REQUIRED_SCORE_CELLS if cells[label].uncached]
    if uncached:
        findings.append(
            finding("formula_without_cached_value", Severity.ERROR, _uncached_message(uncached))
        )
        return None
    if raw_name is None:
        message = (
            f"Name holds the Excel error {name_cell.value}, not a name"
            if name_cell.error
            else "Name is blank"
        )
        findings.append(finding("name_missing", Severity.ERROR, message))
        return None
    score = coerce_int(cells["Score Shot"].value)
    if score is None:
        findings.append(
            finding(
                "score_missing",
                Severity.ERROR,
                "Score Shot is blank or not a whole number",
            )
        )
        return None
    if not 0 <= score <= MAX_SCORE:
        findings.append(
            finding(
                "score_out_of_range",
                Severity.ERROR,
                f"Score Shot {score} is outside 0-{MAX_SCORE}",
            )
        )
        return None
    if event_date is None:
        findings.append(finding("unparseable_date", Severity.ERROR, "Event is not a readable date"))
        return None
    status = _status(cells["Status"], finding, findings)
    gauge_class = _gauge_class(cells["Class"], finding, findings)
    if len(name_key(raw_name).split()) < 2:
        findings.append(
            finding(
                "first_name_only",
                Severity.INFO,
                "Only one name was given; it counts as a new shooter for this "
                "event until an admin merges it",
            )
        )
    return ScoreRow(
        row_number=row,
        raw_name=raw_name,
        score=score,
        event_date=event_date,
        status=status,
        gauge_class=gauge_class,
    )


def _left_blank(label: str, finding: MakeFinding) -> Finding:
    """A kept row's optional cell whose formula has no saved value."""
    return finding(
        "formula_without_cached_value",
        Severity.WARNING,
        f"{label} holds a formula with no saved value, so it was left blank; "
        f"{RESAVE_HINT} to keep it",
    )


def _status(cell: _Cell, finding: MakeFinding, findings: list[Finding]) -> str | None:
    if cell.uncached:
        findings.append(_left_blank("Status", finding))
        return None
    if is_blank(cell.value):
        return None
    text = str(cell.value).strip()
    if text.casefold() in _STATUSES:
        return text.casefold()
    findings.append(
        finding(
            "unknown_status",
            Severity.WARNING,
            f"Status '{text}' is not Member, Guest or Deceased; it was left blank",
        )
    )
    return None


def _gauge_class(cell: _Cell, finding: MakeFinding, findings: list[Finding]) -> str | None:
    if cell.uncached:
        findings.append(_left_blank("Class", finding))
        return None
    if is_blank(cell.value):
        return None
    text = str(cell.value).strip()
    key = "".join(text.split()).casefold()
    if key in _GAUGES:
        return _GAUGES[key]
    if key in _CORRECTED_GAUGES:
        corrected = _CORRECTED_GAUGES[key]
        findings.append(
            finding(
                "gauge_class_corrected",
                Severity.WARNING,
                f"Class '{text}' does not exist; it was read as {corrected}",
            )
        )
        return corrected
    findings.append(
        finding(
            "unknown_gauge_class",
            Severity.WARNING,
            f"Class '{text}' is not a known gauge class; it was kept as written",
        )
    )
    return text


def _parse_attendance_sheet(lw: LoadedWorkbook, findings: list[Finding]) -> list[AttendanceRow]:
    values, formulas = _sheet_pair(lw, ATTENDANCE_SHEET)
    header_row, columns = _find_header(values, ATTENDANCE_HEADER)
    letter = get_column_letter(columns["Date"])
    # "=A48+7": the row above plus N days (the only formula the parser evaluates).
    previous_plus_days = re.compile(rf"=\$?{letter}\$?(\d+)\s*\+\s*(\d+)")
    resolved: dict[int, date] = {}
    rows: list[AttendanceRow] = []
    for row in range(header_row + 1, values.max_row + 1):
        date_cell = _cell(values, formulas, row, columns["Date"])
        event_date, evaluated = _resolve_date(row, date_cell, previous_plus_days, resolved)
        if event_date is not None:
            resolved[row] = event_date
        count_cell = _cell(values, formulas, row, columns["Count"])
        if count_cell.blank:
            continue
        attendance = _attendance_row(
            row,
            event_date=event_date,
            evaluated=evaluated,
            date_cell=date_cell,
            count_cell=count_cell,
            findings=findings,
        )
        if attendance is not None:
            rows.append(attendance)
    return rows


def _resolve_date(
    row: int,
    cell: _Cell,
    previous_plus_days: re.Pattern[str],
    resolved: dict[int, date],
) -> tuple[date | None, bool]:
    if not cell.uncached:
        return coerce_date(cell.value), False
    match = previous_plus_days.fullmatch(str(cell.formula))
    if match is None:
        return None, False
    try:
        source, days = int(match[1]), int(match[2])
        if source != row - 1 or source not in resolved:
            return None, False
        return resolved[source] + timedelta(days=days), True
    except (OverflowError, ValueError):
        # Past 9999-12-31, more days than a timedelta holds, or too many digits.
        return None, False


def _attendance_row(
    row: int,
    *,
    event_date: date | None,
    evaluated: bool,
    date_cell: _Cell,
    count_cell: _Cell,
    findings: list[Finding],
) -> AttendanceRow | None:
    def finding(code: str, severity: Severity, message: str) -> Finding:
        return Finding(
            code,
            severity,
            message,
            sheet=ATTENDANCE_SHEET,
            row=row,
            event_date=event_date,
        )

    uncached: list[str] = []
    if date_cell.uncached and not evaluated:
        uncached.append("Date")
    if count_cell.uncached:
        uncached.append("Count")
    if uncached:
        findings.append(
            finding("formula_without_cached_value", Severity.ERROR, _uncached_message(uncached))
        )
        return None
    if event_date is None:
        findings.append(finding("unparseable_date", Severity.ERROR, "Date is not a readable date"))
        return None
    head_count = coerce_int(count_cell.value)
    if head_count is None or head_count < 0:
        findings.append(
            finding(
                "head_count_invalid",
                Severity.ERROR,
                f"Count '{count_cell.value}' is not a whole number of shooters",
            )
        )
        return None
    if evaluated:
        findings.append(
            finding(
                "date_formula_evaluated",
                Severity.INFO,
                f"Date formula {date_cell.formula} was worked out as {event_date.isoformat()}",
            )
        )
    return AttendanceRow(row_number=row, event_date=event_date, head_count=head_count)
