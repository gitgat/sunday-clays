"""Parse the rolling station-scores workbook (Contract C3).

parse_stations alone decides which sheets and rows are kept, and it emits a
finding for every sheet or row it leaves out. Totals are always recomputed
from the hits; the sheet's own SUM/MEDIAN cells are never read.
"""

from collections import defaultdict
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date

from openpyxl.worksheet.worksheet import Worksheet

from sunday_clays.ingest.names import clean_display_name, name_key
from sunday_clays.ingest.types import (
    Finding,
    Severity,
    StationHitsRow,
    StationLayoutEntry,
    StationSheet,
    StationsParse,
)
from sunday_clays.ingest.workbook import (
    LoadedWorkbook,
    coerce_date,
    coerce_int,
    is_blank,
    is_station_sheet,
    name_text,
    normalize_label,
    worksheets,
)
from sunday_clays.station_label import label_number, parse_label

LABEL_SEARCH_ROWS = 60
LABEL_SEARCH_COLUMNS = 10
# Every row scan touches each station column, and openpyxl creates a cell per
# access, so a crafted wide layout on a tall sheet would sidestep the row
# limit's memory guard. Real tabs have 7 stations.
MAX_STATION_COLUMNS = 30

_NO_LAYOUT = "has no readable STATION # / TARGET COUNT / Name layout"
_TOO_MANY_STATIONS = f"has more than {MAX_STATION_COLUMNS} station columns"
_NAME_IN_STATIONS = "has its Name column inside the station columns"
_DUPLICATE_STATION = "has the same station label in two columns"


@dataclass(frozen=True)
class _Layout:
    entries: tuple[StationLayoutEntry, ...]
    columns: tuple[int, ...]  # sheet column of each entry, same order
    name_column: int
    first_row: int  # first row below the "Name" header


@dataclass(frozen=True)
class _Candidate:
    sheet: StationSheet
    tab_date: date | None


def parse_stations(lw: LoadedWorkbook) -> StationsParse:
    findings: list[Finding] = []
    candidates: list[_Candidate] = []
    for sheet in worksheets(lw.values):
        if is_station_sheet(sheet):
            candidate = _parse_sheet(sheet, findings)
            if candidate is not None:
                candidates.append(candidate)
    kept = _resolve_duplicate_dates(candidates, findings)
    return StationsParse(
        sheets=tuple(sorted(kept, key=lambda s: (s.event_date, s.sheet_name))),
        findings=tuple(findings),
    )


def _parse_sheet(sheet: Worksheet, findings: list[Finding]) -> _Candidate | None:
    title = sheet.title
    layout = _read_layout(sheet)
    if isinstance(layout, str):
        findings.append(
            Finding(
                "station_layout_missing",
                Severity.ERROR,
                f"Sheet '{title}' {layout}; the sheet was left out",
                sheet=title,
            )
        )
        return None
    a2_date = coerce_date(sheet.cell(row=2, column=1).value)
    tab_date = coerce_date(title)
    event_date = a2_date if a2_date is not None else tab_date
    rows = _RowReader(sheet, layout, event_date, findings).read()
    if not rows:
        findings.append(
            Finding(
                "empty_station_sheet",
                Severity.WARNING,
                f"Sheet '{title}' has no shooter rows; it was left out",
                sheet=title,
                event_date=event_date,
            )
        )
        return None
    if event_date is None:
        findings.append(
            Finding(
                "sheet_date_missing",
                Severity.ERROR,
                f"Sheet '{title}' has no readable date in A2 or in its tab name; "
                "the sheet was left out",
                sheet=title,
            )
        )
        return None
    if a2_date is None:
        findings.append(
            Finding(
                "sheet_date_from_name",
                Severity.WARNING,
                "A2 is not a readable date, so the tab name date "
                f"{event_date.isoformat()} was used",
                sheet=title,
                event_date=event_date,
            )
        )
    elif tab_date is not None and tab_date != a2_date:
        findings.append(
            Finding(
                "sheet_name_date_mismatch",
                Severity.WARNING,
                f"Tab name says {tab_date.isoformat()} but A2 says "
                f"{a2_date.isoformat()}; A2 was used",
                sheet=title,
                event_date=a2_date,
            )
        )
    station_sheet = StationSheet(
        sheet_name=title, event_date=event_date, layout=layout.entries, rows=rows
    )
    return _Candidate(sheet=station_sheet, tab_date=tab_date)


def _find_label(sheet: Worksheet, label: str, first_row: int) -> tuple[int, int] | None:
    last_row = min(sheet.max_row, LABEL_SEARCH_ROWS)
    last_column = min(sheet.max_column, LABEL_SEARCH_COLUMNS)
    for row in range(first_row, last_row + 1):
        for column in range(1, last_column + 1):
            if normalize_label(sheet.cell(row, column).value) == label:
                return row, column
    return None


def _read_layout(sheet: Worksheet) -> _Layout | str:
    """The sheet's layout, or why it has none (the finding's message)."""
    station_label = _find_label(sheet, "station #", 1)
    target_label = (
        _find_label(sheet, "target count", station_label[0] + 1)
        if station_label is not None
        else None
    )
    name_label = (
        _find_label(sheet, "name", target_label[0] + 1) if target_label is not None else None
    )
    if station_label is None or target_label is None or name_label is None:
        return _NO_LAYOUT
    entries: list[StationLayoutEntry] = []
    columns: list[int] = []
    seen: set[str] = set()
    # Station columns run right of "STATION #" until the first cell that is neither a number nor
    # a label such as "7A" (the fixture's "Total" in column I).
    for column in range(station_label[1] + 1, sheet.max_column + 1):
        cell = sheet.cell(station_label[0], column).value
        label = parse_label(cell)
        if label is None:
            if coerce_int(cell) is not None:  # a number that is not a station: 0, -1, 100
                return _NO_LAYOUT
            break
        if len(entries) == MAX_STATION_COLUMNS:
            return _TOO_MANY_STATIONS
        if label in seen:
            return _DUPLICATE_STATION
        target_count = coerce_int(sheet.cell(target_label[0], column).value)
        if target_count is None or target_count < 1:
            return _NO_LAYOUT
        seen.add(label)
        entries.append(StationLayoutEntry(label_number(label), target_count, label))
        columns.append(column)
    if not entries:
        return _NO_LAYOUT
    if name_label[1] in columns:
        return _NAME_IN_STATIONS
    return _Layout(
        entries=tuple(entries),
        columns=tuple(columns),
        name_column=name_label[1],
        first_row=name_label[0] + 1,
    )


@dataclass(frozen=True)
class _RowReader:
    sheet: Worksheet
    layout: _Layout
    event_date: date | None
    findings: list[Finding]

    def read(self) -> tuple[StationHitsRow, ...]:
        rows: list[StationHitsRow] = []
        for row in range(self.layout.first_row, self.sheet.max_row + 1):
            raw_name = name_text(self.sheet.cell(row, self.layout.name_column).value)
            cells = [self.sheet.cell(row, column).value for column in self.layout.columns]
            if raw_name is None:
                if not all(is_blank(value) for value in cells):
                    self._add(
                        "station_row_missing_name",
                        Severity.WARNING,
                        "Row has station hits but no name; it was skipped",
                        row,
                    )
                continue
            hits_row = self._hits_row(row, raw_name, cells)
            if hits_row is not None:
                rows.append(hits_row)
        return tuple(rows)

    def _hits_row(self, row: int, raw_name: str, cells: Sequence[object]) -> StationHitsRow | None:
        name = clean_display_name(raw_name)
        entries = self.layout.entries
        if all(is_blank(value) for value in cells):
            self._add("station_row_blank", Severity.INFO, "No station hits", row, name)
            return None
        parsed = [coerce_int(value) for value in cells]
        hits = [value for value in parsed if value is not None]
        if len(hits) != len(parsed):
            missing = [
                entry.label for entry, value in zip(entries, parsed, strict=True) if value is None
            ]
            self._add(
                "hits_missing",
                Severity.ERROR,
                "Hits are blank or not whole numbers at station " + ", ".join(missing),
                row,
                name,
            )
            return None
        out_of_range = [
            f"station {entry.label} has {count} of {entry.target_count}"
            for entry, count in zip(entries, hits, strict=True)
            if not 0 <= count <= entry.target_count
        ]
        if out_of_range:
            self._add(
                "hits_out_of_range",
                Severity.ERROR,
                "Hits outside 0..target count: " + "; ".join(out_of_range),
                row,
                name,
            )
            return None
        if len(name_key(raw_name).split()) < 2:
            self._add(
                "first_name_only",
                Severity.INFO,
                "Only one name was given; it counts as a new shooter for this "
                "event until an admin merges it",
                row,
                name,
            )
        station_hits = tuple(
            (entry.label, count) for entry, count in zip(entries, hits, strict=True)
        )
        return StationHitsRow(row_number=row, raw_name=raw_name, hits=station_hits)

    def _add(
        self,
        code: str,
        severity: Severity,
        message: str,
        row: int,
        name: str | None = None,
    ) -> None:
        self.findings.append(
            Finding(
                code,
                severity,
                message,
                sheet=self.sheet.title,
                row=row,
                event_date=self.event_date,
                name=name,
            )
        )


def _resolve_duplicate_dates(
    candidates: list[_Candidate], findings: list[Finding]
) -> list[StationSheet]:
    """Two kept sheets with one A2 date: keep the one whose tab name agrees.

    If both or neither agree, both are left out, so that week keeps whatever
    an earlier import said about it.
    """
    by_date: dict[date, list[_Candidate]] = defaultdict(list)
    for candidate in candidates:
        by_date[candidate.sheet.event_date].append(candidate)
    kept: list[StationSheet] = []
    for event_date, group in by_date.items():
        matching = [c for c in group if c.tab_date == event_date]
        winner: _Candidate | None = None
        if len(group) == 1:
            winner = group[0]
        elif len(matching) == 1:
            winner = matching[0]
        for candidate in group:
            if candidate is winner:
                kept.append(candidate.sheet)
                continue
            others = ", ".join(
                f"'{other.sheet.sheet_name}'" for other in group if other is not candidate
            )
            findings.append(
                Finding(
                    "duplicate_sheet_date",
                    Severity.ERROR,
                    f"A2 date {event_date.isoformat()} is also used by sheet "
                    f"{others}; this sheet was left out",
                    sheet=candidate.sheet.sheet_name,
                    event_date=event_date,
                )
            )
    return kept
