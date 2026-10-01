"""Safe workbook loading and the cell helpers both parsers share (Contract C3)."""

import io
import re
import warnings
import zipfile
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime, timedelta

import openpyxl
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from sunday_clays.ingest.names import name_key
from sunday_clays.ingest.types import FileKind, ParseError

SCORES_SHEET = "ALL SCORE DETAIL"
SPECIAL_SHEET = "Special Event"  # Plan 17: a special-event workbook's one sheet
UNREADABLE_MESSAGE = "This file could not be read as an Excel workbook"
TOO_LARGE_MESSAGE = "File is too large to process"
UNKNOWN_KIND_MESSAGE = (
    "This doesn't look like a Sunday Clays scores, station or special event workbook"
)
# Archive limits (inclusive). The two openpyxl loads peak at ~29x the
# uncompressed size (measured: a 3.8 MB upload with 20.8 MB of sheet XML and
# 701k numeric cells took RSS from 111 MB to 716 MB), so 10 MiB in total bounds
# one upload at ~300 MB. Fixtures: scores 2.70 MB in total, largest member
# 1.81 MB; stations 0.80 MB in total.
MAX_MEMBERS = 500
MAX_MEMBER_BYTES = 8 * 1024 * 1024
MAX_TOTAL_BYTES = 10 * 1024 * 1024
MAX_COMPRESSION_RATIO = 100
MAX_SHEET_ROWS = 100_000

_EXCEL_EPOCH = date(1899, 12, 30)
_FIRST_SERIAL = 61  # 1900-03-01: lower serials straddle Excel's fake 1900-02-29
_LAST_SERIAL = 2_958_465  # 9999-12-31
_US_DATE = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{4})")
# A time suffix (" " or "T", then a digit: seconds, fractions, "Z", offsets) is
# ignored; free text after the date is not a date.
_ISO_DATE = re.compile(r"(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T]\d.*)?")
_TAB_DATE = re.compile(r"(\d{1,2})\s+(\d{1,2})\s+(\d{2})")


@dataclass(frozen=True)
class LoadedWorkbook:
    values: Workbook  # cached cell values (data_only=True)
    formulas: Workbook  # formula text where a cell holds a formula


def load_workbook_bytes(data: bytes) -> LoadedWorkbook:
    """Load an upload twice: cached values and formula text.

    Raises ParseError with a user-facing message for anything unusable.
    """
    _check_archive(data)
    values = _load(data, data_only=True)
    # openpyxl creates a cell on every access, so a stray value far down a
    # sheet would make the row loops allocate millions of empty cells.
    if any(sheet.max_row > MAX_SHEET_ROWS for sheet in worksheets(values)):
        raise ParseError(TOO_LARGE_MESSAGE)
    return LoadedWorkbook(values=values, formulas=_load(data, data_only=False))


def _check_archive(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
    except Exception as exc:
        raise ParseError(UNREADABLE_MESSAGE) from exc
    if _archive_too_large(members):
        raise ParseError(TOO_LARGE_MESSAGE)


def _archive_too_large(members: Sequence[zipfile.ZipInfo]) -> bool:
    """Whether the declared sizes break a limit; every limit is inclusive.

    A member with ``compress_size`` 0 is treated as 1 compressed byte, so the
    ratio check never divides by zero.
    """
    return (
        len(members) > MAX_MEMBERS
        or sum(member.file_size for member in members) > MAX_TOTAL_BYTES
        or any(member.file_size > MAX_MEMBER_BYTES for member in members)
        or any(
            member.file_size > MAX_COMPRESSION_RATIO * max(member.compress_size, 1)
            for member in members
        )
    )


def _load(data: bytes, *, data_only: bool) -> Workbook:
    try:
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore",
                message="Data Validation extension is not supported",
                category=UserWarning,
                module="openpyxl",
            )
            return openpyxl.load_workbook(io.BytesIO(data), data_only=data_only)
    except Exception as exc:
        # Note for test authors: pytest runs with filterwarnings=error, so any
        # openpyxl warning other than the Data Validation one is raised inside
        # load_workbook and lands here. The test then sees UNREADABLE for a
        # file that production (where warnings stay warnings) loads fine.
        raise ParseError(UNREADABLE_MESSAGE) from exc


def worksheets(workbook: Workbook) -> list[Worksheet]:
    """The workbook's grid sheets, in tab order."""
    return [sheet for sheet in workbook.worksheets if isinstance(sheet, Worksheet)]


def normalize_label(value: object) -> str | None:
    """casefold(strip(collapse whitespace)) for text cells, else None."""
    if not isinstance(value, str):
        return None
    return " ".join(value.split()).casefold()


def is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def name_text(value: object) -> str | None:
    """The cell as a name, or None when it is blank or has no letters/digits."""
    if value is None:
        return None
    text = value if isinstance(value, str) else str(value)
    return text if name_key(text) else None


def coerce_int(value: object) -> int | None:
    """Whole numbers from int, integral float or numeric text; never bool."""
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, str):
        try:
            value = float(value.strip())
        except ValueError:
            return None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return None


def coerce_date(value: object) -> date | None:
    """datetime, date, Excel serial, M/D/YYYY, YYYY-MM-DD or tab-style "9 13 26"."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        if not _FIRST_SERIAL <= value <= _LAST_SERIAL:
            return None
        return _EXCEL_EPOCH + timedelta(days=int(value))
    if isinstance(value, str):
        return _date_from_text(value.strip())
    return None


def _date_from_text(text: str) -> date | None:
    if match := _US_DATE.fullmatch(text):
        month, day, year = (int(part) for part in match.groups())
    elif match := _ISO_DATE.fullmatch(text):
        year, month, day = (int(part) for part in match.groups())
    elif match := _TAB_DATE.fullmatch(text):
        month, day, short_year = (int(part) for part in match.groups())
        year = 2000 + short_year
    else:
        return None
    try:
        return date(year, month, day)
    except ValueError:
        return None


def find_sheet(workbook: Workbook, title: str) -> Worksheet | None:
    """The sheet whose title matches ``title`` ignoring case and spacing."""
    wanted = normalize_label(title)
    for sheet in worksheets(workbook):
        if normalize_label(sheet.title) == wanted:
            return sheet
    return None


def is_station_sheet(sheet: Worksheet) -> bool:
    return normalize_label(sheet.cell(row=1, column=1).value) == "event date"


def detect_kind(lw: LoadedWorkbook) -> FileKind:
    """Scores sheet first, then the special-event sheet, then station tabs."""
    if find_sheet(lw.values, SCORES_SHEET) is not None:
        return FileKind.SCORES
    if find_sheet(lw.values, SPECIAL_SHEET) is not None:
        return FileKind.SPECIAL
    if any(is_station_sheet(sheet) for sheet in worksheets(lw.values)):
        return FileKind.STATIONS
    raise ParseError(UNKNOWN_KIND_MESSAGE)
